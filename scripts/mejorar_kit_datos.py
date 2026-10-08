"""Mejoras de uso del Kit Ley 21.719 (octubre de 2026).

Motivo: un comprador (y el propio dueño del sitio) no entendió cómo se completa el Excel RAT. La guía del
Excel decía "cambia la fila 3" sin explicar qué hay que llenar, qué columnas significan ni cuándo borrar
filas, y no mencionaba la fecha. Este script:

- Reescribe la pestaña "Guía Rápida" del RAT (pasos, significado de cada columna, qué cambiar en cada fila
  pre-llenada), marca en amarillo las 3 celdas que se completan, agrega la pestaña "Fila nueva" con un ejemplo
  y pasa el Excel a los colores de la marca (verde #00382E).
- Corrige el manual (Word y HTML): la fila RAT-05 no es "Ventas web" (es RAT-04), la fila 3 lleva también la
  fecha, y se remite a la Guía Rápida.
- Documento 7: la columna SÍ del test no tenía casilla; se agrega, y se indica a quién se designa si la
  respuesta a la pregunta 1 es SÍ.
- Documentos 2, 5 y 7: "Conforme a la Ley" pasa a "Basada en la Ley" (no prometemos conformidad, AGENTS.md 2.1).
- Documento 3 (DPA): cláusula 3 más clara (24 h del encargado al responsable; 72 h es meta interna del responsable).

Es idempotente. Orden de uso:
    python scripts/mejorar_kit_datos.py                  (edita los archivos sueltos de api/assets/kit_datos_files)
    python scripts/render_manual_datos_pdf.py            (regenera el PDF del manual desde su HTML)
    python scripts/mejorar_kit_datos.py --empaquetar     (reescribe el zip y el módulo base64 de api/assets)
"""
import base64
import io
import math
import os
import sys
import zipfile

import docx
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docs_layout as L  # noqa: E402
from fix_kit_content import _paragraphs, _replace_in_paragraph  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(BASE, 'api', 'assets')
SUELTOS = os.path.join(ASSETS, 'kit_datos_files')
ZIP = os.path.join(ASSETS, 'Kit_Ley_21719_Proteccion_Datos_Pyme_2026.zip')
B64 = os.path.join(ASSETS, 'kit-datos-base64.js')
ORDEN = [
    '0_MANUAL_DE_USO_GUIA_RAPIDA_PYMES.pdf',
    '0_MANUAL_DE_USO_GUIA_RAPIDA_PYMES.docx',
    '1_Anexo_Laboral_Datos_Personales_Ley_21719.docx',
    '2_Politica_Privacidad_Web_y_Pyme_Ley_21719.docx',
    '3_Clausula_DPA_Proveedores_Encargados_Ley_21719.docx',
    '4_Registro_Actividades_Tratamiento_RAT_Ley_21719.xlsx',
    '5_Protocolo_Brechas_Seguridad_72h_Ley_21719.docx',
    '6_Formulario_Solicitud_Derechos_ARCOP.docx',
    '7_Guia_Autodiagnostico_DPO_Delegado_Proteccion_Datos_Pyme.docx',
]
MANUAL_HTML = '0_Manual_Instrucciones_Paso_a_Paso_Ley_21719.html'
MANUAL_DOCX = '0_MANUAL_DE_USO_GUIA_RAPIDA_PYMES.docx'
RAT = '4_Registro_Actividades_Tratamiento_RAT_Ley_21719.xlsx'

# ---------------------------------------------------------------------------------------------- Word
OLD_P2 = ('Abre el Documento 4 (Excel RAT). En la Fila 3, cambia [NOMBRE DE TU EMPRESA] y [RUT EMPRESA] por tus datos '
          'reales y guárdalo en tu computador.')
NEW_P2 = ('Abre el Documento 4 (Excel RAT) y lee primero su pestaña «Guía Rápida» (3 minutos). Luego, en la pestaña '
          '«Registro RAT», completa las 3 celdas amarillas de la fila 3 (razón social, RUT y mes y año), revisa las 6 filas '
          'ya rellenas (borra las que tu empresa no usa) y guárdalo en tu computador.')

WORD = {
    MANUAL_DOCX: [
        (OLD_P2, NEW_P2),
        ('mantén activa la fila RAT-05 (Ventas web) y nombra a los couriers en Columna H.',
         'mantén activa la fila RAT-04 (Ventas y boletas) y nombra tu pasarela de pago y a los couriers en la Columna H.'),
        ('En el RAT añade la fila RAT-07: Rastreo GPS.',
         'En el RAT añade la fila RAT-07: Rastreo GPS (usa la pestaña «Fila nueva» del Excel como modelo).'),
        ('Planilla completada con el Nombre y RUT en la Fila 3, archivada en formato Excel y PDF en el computador de administración.',
         'Planilla con razón social, RUT y fecha completados (fila 3), filas revisadas (borradas las que no aplican) y '
         'archivada en formato Excel y PDF en el computador de administración.'),
    ],
    '2_Politica_Privacidad_Web_y_Pyme_Ley_21719.docx': [
        ('Conforme a la Ley N° 21.719 sobre Protección y Tratamiento de Datos Personales de la República de Chile',
         'Basada en la Ley N° 21.719 sobre Protección y Tratamiento de Datos Personales de la República de Chile'),
    ],
    '5_Protocolo_Brechas_Seguridad_72h_Ley_21719.docx': [
        ('(Conforme al Artículo 14 sexies de la Ley N° 19.628 reformada por Ley N° 21.719 - Estándar Operativo 72 Horas)',
         '(Basado en el Artículo 14 sexies de la Ley N° 19.628 reformada por Ley N° 21.719 - Estándar Operativo 72 Horas)'),
    ],
    '7_Guia_Autodiagnostico_DPO_Delegado_Proteccion_Datos_Pyme.docx': [
        ('Conforme a la Ley N° 21.719 que reforma la Ley N° 19.628', 'Basada en la Ley N° 21.719 que reforma la Ley N° 19.628'),
        ('Responsable interno designado: [NOMBRE Y CARGO].',
         'Responsable interno designado: [NOMBRE Y CARGO]. Si respondió SÍ a la pregunta 1, el Delegado de Protección de '
         'Datos designado es: [NOMBRE Y CARGO DEL DELEGADO].'),
    ],
    '3_Clausula_DPA_Proveedores_Encargados_Ley_21719.docx': [
        ('datos afectados y medidas correctivas inmediatas, a fin de que el Responsable pueda cumplir con la obligación de '
         'notificar a la Agencia de Protección de Datos Personales sin dilaciones indebidas (conforme al Art. 14 sexies de la '
         'Ley N° 19.628 reformada), adoptando el estándar operativo y mejor práctica internacional recomendada de un plazo '
         'máximo de 72 horas.',
         'datos afectados y medidas correctivas inmediatas. Este plazo permite que el Responsable cumpla su obligación de '
         'reportar a la Agencia de Protección de Datos Personales sin dilaciones indebidas (Art. 14 sexies de la Ley N° 19.628 '
         'reformada), para lo cual usa como meta interna un máximo de 72 horas desde que detecta el incidente.'),
    ],
}


def arreglar_word(nombre, reglas):
    ruta = os.path.join(SUELTOS, nombre)
    d = docx.Document(ruta)
    hechos = 0
    for old, new in reglas:
        if any(new in p.text for p in _paragraphs(d)):
            continue  # ya aplicado (el texto nuevo puede contener al viejo: no se repite)
        n = sum(_replace_in_paragraph(p, old, new) for p in _paragraphs(d))
        if n:
            hechos += n
        else:
            print('   FALTA el texto:', old[:80])
    if nombre.startswith('7_Guia'):
        hechos += casillas_si_doc7(d)
    if hechos:
        if not nombre.startswith('0_MANUAL'):
            L.highlight_fields(d)
        d.save(ruta)
    print('  %-66s %d cambio(s)' % (nombre, hechos))


def casillas_si_doc7(d):
    """La tabla del test tenía casilla solo en la columna NO: se agrega en la columna SÍ."""
    n = 0
    for t in d.tables:
        if len(t.columns) != 3 or t.rows[0].cells[1].text.strip().upper() != 'SÍ':
            continue
        for fila in t.rows[1:]:
            celda = fila.cells[1]
            if celda.text.strip() == '' and fila.cells[2].text.strip():
                modelo = fila.cells[2].paragraphs[0]
                p = celda.paragraphs[0]
                r = p.add_run('[   ]')
                if modelo.runs:
                    mr = modelo.runs[0]
                    r.font.name, r.font.size, r.font.bold = mr.font.name, mr.font.size, mr.font.bold
                p.alignment = modelo.alignment
                n += 1
    return n


# ---------------------------------------------------------------------------------------------- HTML del manual
HTML = [
    ('Abre el <strong>Documento 4 (Excel RAT)</strong>, reemplaza en la fila 3 el Nombre y RUT de tu empresa, y guárdalo.',  # noqa: E501
     'Abre el <strong>Documento 4 (Excel RAT)</strong>: lee su «Guía Rápida», completa la fila 3 (nombre, RUT y fecha) y borra las '
     'filas que no uses.'),
    ('<div class="doc-action">Reemplazar Nombre y RUT en Fila 3 y guardar</div>',
     '<div class="doc-action">Completar fila 3, revisar filas y guardar</div>'),
    ('mantén activa la fila <code>RAT-05</code> (Ventas web). Si usas Starken o Chilexpress, puedes mencionarlos en la Columna H '
     'como encargados de entrega.',
     'mantén activa la fila <code>RAT-04</code> (Ventas y boletas). Anota tu pasarela de pago y, si usas Starken o Chilexpress, '
     'menciónalos en la Columna H como encargados de entrega.'),
    ('indicando su finalidad (seguridad y control de ruta) y que los choferes fueron informados.',
     'indicando su finalidad (seguridad y control de ruta) y que los choferes fueron informados (la pestaña «Fila nueva» del Excel '
     'trae este ejemplo).'),
    ('Planilla completada con el Nombre y RUT de la empresa en la Fila 3, guardada en el computador de administración en formato '
     'Excel y PDF.',
     'Planilla con razón social, RUT y fecha completados (fila 3), filas revisadas (borradas las que no aplican), guardada en el '
     'computador de administración en formato Excel y PDF.'),
]


def arreglar_html():
    ruta = os.path.join(SUELTOS, MANUAL_HTML)
    t = open(ruta, 'r', encoding='utf-8', newline='').read()
    crlf = '\r\n' in t
    c = t.replace('\r\n', '\n')
    hechos = 0
    for old, new in HTML:
        if old in c:
            c = c.replace(old, new)
            hechos += 1
        elif new not in c:
            print('   FALTA el texto del HTML:', old[:80])
    if hechos:
        open(ruta, 'w', encoding='utf-8', newline='').write(c.replace('\n', '\r\n') if crlf else c)
    print('  %-66s %d cambio(s)' % (MANUAL_HTML, hechos))


# ---------------------------------------------------------------------------------------------- Excel RAT
VERDE, VERDE_HONDO, VERDE_TENUE = '00382E', '00261F', 'ECF5F2'
AMARILLO = 'FFF3B0'
PIZARRA, TEXTO, CLARO = '334155', '0F172A', 'F1F5F9'
FINO = Side(style='thin', color='CBD5E1')
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)

# colores del archivo original -> colores de la marca
MAPA = {'000284C7': VERDE, '000F172A': None, '001E293B': VERDE_HONDO, '00F0F9FF': VERDE_TENUE, '00059669': '00047857'}

TABLA_COLUMNAS = [
    ('Código RAT', 'Un número correlativo para ubicar la fila.', 'RAT-01, RAT-02… Si agregas una fila, sigue la numeración: RAT-07.'),
    ('Área / Proceso', 'La parte de tu empresa que usa esos datos.', 'Recursos Humanos, Ventas, Bodega, Administración.'),
    ('Categoría de Datos Personales', 'Qué datos concretos guardas (los campos o papeles).', 'Nombre, RUT, correo, teléfono, dirección, cuenta bancaria.'),
    ('Titulares Afectados', 'De quién son esos datos.', 'Trabajadores, clientes, proveedores, visitantes.'),
    ('Finalidad del Tratamiento', 'Para qué los usas, en una frase.', 'Emitir boletas y despachar pedidos.'),
    ('Base Legal de Licitud', 'Por qué puedes usarlos. Las más comunes: contrato, obligación legal, consentimiento, interés legítimo.',
     'Ejecución del contrato de compraventa y obligación tributaria.'),
    ('¿Datos Sensibles?', 'SÍ si son de salud, biometría (huella, rostro) u otros muy personales (ideología, religión, afiliación '
     'sindical, vida sexual, origen étnico). Si trabajas con datos de menores de edad, anótalo también: la ley les da protección especial. '
     'Si no, NO.', 'Licencias médicas = SÍ (Salud). Boletas de venta = NO.'),
    ('Destinatarios / Encargados Externos', 'Con quién compartes los datos: instituciones y proveedores.', 'Banco, Previred, SII, Transbank o Flow, courier, contador.'),
    ('Plazo de Conservación', 'Cuánto tiempo los guardas antes de borrarlos. Si no sabes, deja el que viene.', '6 años (respaldo tributario).'),
    ('Medidas de Seguridad Aplicadas', 'Cómo proteges los datos en la práctica. Escribe solo lo que de verdad haces.',
     'Claves individuales, archivador con llave, respaldo semanal, antivirus.'),
]

TABLA_FILAS = [
    ('RAT-01 · Sueldos y RR.HH.', 'Cambia «Banco pagador» y «Software de Nómina» por el banco y el programa que usas de verdad (Buk, Talana, una planilla Excel…). '
     'Revisa «Medidas de seguridad»: deja solo las que aplicas.'),
    ('RAT-02 · Licencias médicas', 'Déjala si tienes trabajadores. Revisa que quien las ve sea solo RR.HH. (o quien tramita las licencias) y ajusta la mutual o Isapre si corresponde.'),
    ('RAT-03 · Reloj control con huella o rostro', 'Si NO usas huella ni rostro, bórrala (la tarjeta, la clave o el libro no son biométricos). Si la usas: escribe el nombre del proveedor del reloj.'),
    ('RAT-04 · Ventas y boletas', 'Anota tu pasarela de pago real (Transbank, Flow, Mercado Pago…) y tus couriers. Si vendes solo a empresas, ajusta «Titulares».'),
    ('RAT-05 · Marketing y WhatsApp', 'Si no envías promociones, bórrala. Si las envías: escribe tu herramienta (Mailchimp, WhatsApp Business…) y confirma que cada mensaje deja darse de baja.'),
    ('RAT-06 · Cámaras de seguridad', 'Si no tienes cámaras, bórrala. Si las tienes: cambia «30 días» por el tiempo real que graba tu equipo.'),
]

PASOS = [
    ('Paso 1', 'Abre la pestaña «Registro RAT» (abajo, junto a esta).'),
    ('Paso 2', 'Completa las 3 celdas amarillas de la fila 3:\n• C3: razón social de tu empresa\n• F3: RUT de la empresa\n• I3: mes y año de hoy (por ejemplo «Octubre 2026»).'),
    ('Paso 3', 'Recorre las filas RAT-01 a RAT-06 y pregúntate en cada una: «¿mi empresa hace esto?»\n'
               '• SÍ → déjala y corrige lo que sea distinto en tu caso (mira la sección 5).\n'
               '• NO → bórrala: clic derecho sobre el número de la fila → Eliminar.\n'
               '• Si hay algo que tu empresa hace y no aparece (por ejemplo, GPS en las camionetas): copia la fila de la pestaña «Fila nueva» a la tabla y cámbiale los textos.'),
    ('Paso 4', 'Revisa que cada fila diga solo lo que es verdad en tu empresa, sobre todo «Medidas de seguridad»: no dejes escrito doble factor, cifrado ni certificaciones que no tengas.'),
    ('Paso 5', 'Guarda el archivo (y, si quieres, también un PDF o una impresión) en la carpeta de administración. Vuelve a abrirlo y cambia la fecha cuando algo cambie: un software nuevo, un proveedor nuevo, una cámara nueva.'),
]

AVISO = ('Modelo de referencia basado en la normativa vigente a octubre de 2026; no constituye asesoría legal. Si tu caso tiene '
         'particularidades, revísalo con un abogado antes de usarlo. La Ley 21.719 rige desde el 1 de diciembre de 2026.')


def _lineas(texto, ancho):
    n = 0
    for tramo in str(texto).split('\n'):
        n += max(1, math.ceil(len(tramo) / max(ancho, 1)))
    return n


def _alto(texto, ancho_cols, pt=10):
    # ~1,15 caracteres por unidad de ancho a 10 pt; 13 pt por línea
    lineas = _lineas(texto, ancho_cols * (1.15 if pt <= 10 else 1.0))
    return max(20, lineas * (pt + 3.5) + 8)


def recolor(ws):
    for fila in ws.iter_rows():
        for c in fila:
            if c.font and c.font.color is not None and c.font.color.type == 'rgb' and c.font.color.rgb in MAPA:
                nuevo = MAPA[c.font.color.rgb]
                if nuevo:
                    c.font = Font(name=c.font.name, size=c.font.sz, bold=c.font.b, italic=c.font.i, color=nuevo)
            if c.fill and c.fill.fill_type and c.fill.fgColor.rgb in MAPA:
                nuevo = MAPA[c.fill.fgColor.rgb]
                if c.fill.fgColor.rgb == '000F172A':
                    nuevo = VERDE
                if nuevo:
                    c.fill = PatternFill('solid', fgColor=nuevo)


def _titulo(ws, rango, texto, fondo, color, size, bold, alto):
    ws.merge_cells(rango)
    c = ws[rango.split(':')[0]]
    c.value = texto
    c.font = Font(name='Calibri', size=size, bold=bold, color=color)
    c.fill = PatternFill('solid', fgColor=fondo)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[c.row].height = alto


def construir_guia(wb):
    nombre = '📌 Guía Rápida'
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre, 0)
    anchos = {'A': 24, 'B': 15, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 15, 'H': 15}
    for k, v in anchos.items():
        ws.column_dimensions[k].width = v
    ws.sheet_view.showGridLines = False
    total = sum(anchos.values())
    r = 1
    _titulo(ws, f'A{r}:H{r}', 'GUÍA RÁPIDA: CÓMO COMPLETAR TU REGISTRO RAT', VERDE, 'FFFFFFFF', 14, True, 36); r += 1
    _titulo(ws, f'A{r}:H{r}', 'Léela una vez (3 minutos). Después trabaja en la pestaña «Registro RAT».', VERDE_HONDO, 'FFE2E8F0', 10, False, 22); r += 2

    def seccion(texto):
        nonlocal r
        ws.merge_cells(f'A{r}:H{r}')
        c = ws[f'A{r}']
        c.value = texto
        c.font = Font(name='Calibri', size=11, bold=True, color=VERDE)
        c.fill = PatternFill('solid', fgColor=CLARO)
        c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        c.border = BORDE
        ws.row_dimensions[r].height = 24
        r += 1

    def parrafo(texto):
        nonlocal r
        ws.merge_cells(f'A{r}:H{r}')
        c = ws[f'A{r}']
        c.value = texto
        c.font = Font(name='Calibri', size=10, color=PIZARRA)
        c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True, indent=1)
        c.border = BORDE
        ws.row_dimensions[r].height = _alto(texto, total - 2)
        r += 1

    seccion('1. ¿QUÉ ES ESTE ARCHIVO?')
    parrafo('El RAT (Registro de Actividades de Tratamiento) es un inventario: una tabla donde dejas por escrito qué datos de personas '
            'maneja tu empresa (de trabajadores, clientes, proveedores), para qué los usas, con quién los compartes y cuánto tiempo los guardas.\n'
            'Es como el libro de compras y ventas, pero de datos personales. Sirve para mostrar orden si la Agencia de Protección de Datos '
            'Personales te pregunta o recibes un reclamo.')
    r += 1
    seccion('2. ¿LO TENGO QUE ENVIAR A ALGUIEN?')
    parrafo('No. No se envía a ninguna institución. Lo guardas tú, en la carpeta de administración, y lo muestras solo si la Agencia de '
            'Protección de Datos o la Dirección del Trabajo te lo piden.\n'
            'Casi todo viene ya escrito con los casos típicos de una pyme chilena: tú completas 3 celdas y ajustas lo que sea distinto.')
    r += 1
    seccion('3. LO QUE TIENES QUE HACER (5 PASOS, ~15 MINUTOS)')
    for paso, texto in PASOS:
        ws.merge_cells(f'B{r}:H{r}')
        a, b = ws[f'A{r}'], ws[f'B{r}']
        a.value, b.value = paso, texto
        a.font = Font(name='Calibri', size=10, bold=True, color='FFFFFFFF')
        a.fill = PatternFill('solid', fgColor=VERDE)
        a.alignment = Alignment(horizontal='center', vertical='center')
        b.font = Font(name='Calibri', size=10, color=PIZARRA)
        b.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True, indent=1)
        a.border = b.border = BORDE
        for col in 'CDEFGH':
            ws[f'{col}{r}'].border = BORDE
        ws.row_dimensions[r].height = _alto(texto, total - 24 - 2)
        r += 1
    r += 1
    seccion('4. QUÉ SIGNIFICA CADA COLUMNA DE LA TABLA')
    for a, b, c in [('Columna', 'Qué escribir', 'Ejemplo')]:
        ws.merge_cells(f'B{r}:E{r}'); ws.merge_cells(f'F{r}:H{r}')
        for ref, v in ((f'A{r}', a), (f'B{r}', b), (f'F{r}', c)):
            ws[ref].value = v
            ws[ref].font = Font(name='Calibri', size=10, bold=True, color='FFFFFFFF')
            ws[ref].fill = PatternFill('solid', fgColor=VERDE)
            ws[ref].alignment = Alignment(horizontal='center', vertical='center')
        for col in 'ABCDEFGH':
            ws[f'{col}{r}'].border = BORDE
            ws[f'{col}{r}'].fill = PatternFill('solid', fgColor=VERDE)
        ws.row_dimensions[r].height = 22
        r += 1
    for i, (col, que, ej) in enumerate(TABLA_COLUMNAS):
        ws.merge_cells(f'B{r}:E{r}'); ws.merge_cells(f'F{r}:H{r}')
        fondo = 'F8FAFC' if i % 2 == 0 else 'FFFFFF'
        for ref, v, b in ((f'A{r}', col, True), (f'B{r}', que, False), (f'F{r}', ej, False)):
            ws[ref].value = v
            ws[ref].font = Font(name='Calibri', size=10, bold=b, color=TEXTO if b else PIZARRA)
            ws[ref].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
        for c_ in 'ABCDEFGH':
            ws[f'{c_}{r}'].border = BORDE
            ws[f'{c_}{r}'].fill = PatternFill('solid', fgColor=fondo)
        ws.row_dimensions[r].height = max(_alto(que, 15 * 4 - 2), _alto(ej, 15 * 3 - 2), _alto(col, 22))
        r += 1
    r += 1
    ws.row_breaks.append(Break(id=r - 1))
    seccion('5. QUÉ CAMBIAR EN CADA FILA YA RELLENA')
    for i, (fila, texto) in enumerate(TABLA_FILAS):
        ws.merge_cells(f'B{r}:H{r}')
        fondo = 'F8FAFC' if i % 2 == 0 else 'FFFFFF'
        ws[f'A{r}'].value, ws[f'B{r}'].value = fila, texto
        ws[f'A{r}'].font = Font(name='Calibri', size=10, bold=True, color=VERDE)
        ws[f'B{r}'].font = Font(name='Calibri', size=10, color=PIZARRA)
        ws[f'A{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
        ws[f'B{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
        for c_ in 'ABCDEFGH':
            ws[f'{c_}{r}'].border = BORDE
            ws[f'{c_}{r}'].fill = PatternFill('solid', fgColor=fondo)
        ws.row_dimensions[r].height = max(_alto(texto, total - 24 - 2), _alto(fila, 22))
        r += 1
    r += 1
    seccion('6. SI AÚN NO TIENES EMPRESA')
    parrafo('Puedes usar el archivo como ejemplo ya resuelto de cómo se ve un RAT: no hay nada que enviar ni registrar. '
            'Cuando tengas la empresa, sigue los pasos de la sección 3.')
    r += 1
    ws.merge_cells(f'A{r}:H{r}')
    ws[f'A{r}'].value = AVISO
    ws[f'A{r}'].font = Font(name='Calibri', size=9, italic=True, color='475569')
    ws[f'A{r}'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True, indent=1)
    ws.row_dimensions[r].height = _alto(AVISO, total - 2, 9)
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return ws


def preparar_registro(ws):
    recolor(ws)
    # celdas que se completan: amarillas
    for ref in ('C3', 'F3', 'I3'):
        c = ws[ref]
        c.fill = PatternFill('solid', fgColor=AMARILLO)
        c.font = Font(name='Calibri', size=10, bold=True, color=TEXTO)
    ws.row_dimensions[4].height = 22
    if not ws.merged_cells.ranges or 'A4:J4' not in [str(m) for m in ws.merged_cells.ranges]:
        ws.merge_cells('A4:J4')
    n = ws['A4']
    n.value = ('Celdas amarillas = complétalas. Revisa cada fila (la pestaña «Guía Rápida» explica cómo): déjala si tu empresa lo hace, '
               'ajusta lo distinto y borra las filas que no aplican.')
    n.font = Font(name='Calibri', size=9.5, italic=True, color='475569')
    n.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.freeze_panes = 'A6'


def construir_fila_nueva(wb, registro):
    nombre = '➕ Fila nueva'
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre)
    ws.sheet_view.showGridLines = False
    for col in 'ABCDEFGHIJ':
        ws.column_dimensions[col].width = registro.column_dimensions[col].width
    _titulo(ws, 'A1:J1', 'FILA NUEVA: MODELO PARA AGREGAR UNA ACTIVIDAD QUE NO ESTÁ EN LA TABLA', VERDE, 'FFFFFFFF', 13, True, 32)
    ws.merge_cells('A2:J2')
    ws['A2'].value = ('Copia la fila de abajo, pégala en la primera fila vacía de «Registro RAT» (debajo de la última) y cambia los textos. '
                      'Sigue la numeración: si la última es RAT-06, esta es RAT-07. El ejemplo es de una empresa de transporte con GPS en sus camiones; '
                      'bórralo o reemplázalo por tu caso.')
    ws['A2'].font = Font(name='Calibri', size=9.5, italic=True, color='475569')
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
    ws.row_dimensions[2].height = 42
    for col in 'ABCDEFGHIJ':
        h = registro[f'{col}5']
        c = ws[f'{col}4']
        c.value = h.value
        c.font = Font(name=h.font.name, size=h.font.sz, bold=h.font.b, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor=VERDE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDE
    ws.row_dimensions[4].height = 32
    ejemplo = ['RAT-07', 'Operaciones y transporte', 'Ubicación GPS de camiones y furgones, nombre del chofer asignado',
               'Choferes', 'Seguridad de la carga, control de ruta y atención de emergencias en carretera',
               'Interés legítimo del responsable en la seguridad de sus bienes; choferes informados por escrito (Anexo laboral)',
               'NO', 'Proveedor del sistema GPS (con contrato de tratamiento de datos firmado)',
               'Mientras dure el contrato del chofer; el historial de rutas, 12 meses',
               'Acceso con clave personal solo para jefatura de operaciones; el GPS se usa en horario de trabajo']
    vacia = ['RAT-08', '[Área o proceso]', '[Qué datos guardas]', '[De quién son]', '[Para qué los usas]',
             '[Por qué puedes usarlos]', '[SÍ / NO]', '[Con quién los compartes]', '[Cuánto tiempo los guardas]',
             '[Cómo los proteges, solo lo que haces]']
    for fila, valores, amarillo in ((5, ejemplo, False), (6, vacia, True)):
        for col, v in zip('ABCDEFGHIJ', valores):
            c = ws[f'{col}{fila}']
            c.value = v
            c.font = Font(name='Consolas' if col == 'A' else 'Calibri', size=9.5 if col == 'A' else 9, bold=(col == 'A'),
                          color=VERDE if col == 'A' else TEXTO)
            c.fill = PatternFill('solid', fgColor=AMARILLO if amarillo else 'F8FAFC')
            c.alignment = Alignment(horizontal='center' if col in 'AB' else 'left', vertical='center', wrap_text=True)
            c.border = BORDE
        ws.row_dimensions[fila].height = 58
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def arreglar_excel():
    ruta = os.path.join(SUELTOS, RAT)
    wb = openpyxl.load_workbook(ruta)
    registro = next((ws for ws in wb.worksheets if 'Registro RAT' in ws.title), None)
    if registro is None:
        raise SystemExit('No se encontró la pestaña «Registro RAT»')
    preparar_registro(registro)
    construir_guia(wb)
    construir_fila_nueva(wb, registro)
    # orden: Guía, Registro, Fila nueva; abre en la Guía
    wb._sheets = [wb[n] for n in ('📌 Guía Rápida', registro.title, '➕ Fila nueva')]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = (ws.title == '📌 Guía Rápida')
    wb.properties.title = 'Registro de Actividades de Tratamiento (RAT) - Kit Ley 21.719'
    wb.save(ruta)
    print('  %-66s reescrito (Guía Rápida, celdas amarillas, Fila nueva)' % RAT)


# ---------------------------------------------------------------------------------------------- empaquetado
def empaquetar():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for nombre in ORDEN:
            ruta = os.path.join(SUELTOS, nombre)
            zf.write(ruta, arcname=nombre)
            print('  zip <-', nombre, os.path.getsize(ruta), 'bytes')
    datos = buf.getvalue()
    open(ZIP, 'wb').write(datos)
    with open(B64, 'w', encoding='utf-8', newline='\n') as f:
        f.write('module.exports = "%s";\n' % base64.b64encode(datos).decode())
    print('Kit_Ley_21719_Proteccion_Datos_Pyme_2026.zip', len(datos), 'bytes; kit-datos-base64.js actualizado')


def main():
    if '--empaquetar' in sys.argv:
        empaquetar()
        return
    print('Word:')
    for nombre, reglas in WORD.items():
        arreglar_word(nombre, reglas)
    print('HTML del manual:')
    arreglar_html()
    print('Excel:')
    arreglar_excel()
    print('\nSiguiente: python scripts/render_manual_datos_pdf.py  y luego  python scripts/mejorar_kit_datos.py --empaquetar')


if __name__ == '__main__':
    main()
