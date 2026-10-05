"""Correcciones de contenido de los kits de pago (octubre de 2026).

Revisión legal de los Word y planillas del Kit Blindaje Laboral y del Kit Ley 21.719: promesas que no se
pueden cumplir ("evita multas", "blindaje", "aprobación DT"), cifras de multas exageradas para pymes,
artículos mal citados y errores de fondo (horario que no suma 42 horas, reparto de propinas por el
empleador, plazo de propinas, Art. 22, delegado de protección de datos, etc.).

Fuentes verificadas en el Código del Trabajo (texto vigente a septiembre de 2026): Arts. 9, 9 bis, 22,
38, 38 bis, 64, 153, 154 bis, 154 ter, 152 quáter J y O, 506, 506 ter y 511. Ley 19.628 modificada
por la Ley 21.719 (Art. 11: 30 días corridos prorrogables una vez; Art. 14 sexies: brechas "sin
dilaciones indebidas", sin plazo en horas; delegado de protección de datos voluntario).

Es idempotente: si un texto ya fue corregido, no se vuelve a tocar. Al final reescribe los zip y los
módulos base64 de api/assets. Los manuales PDF se regeneran aparte desde su HTML
(scripts/render_kit_manuals.mjs).

Uso: python scripts/fix_kit_content.py
"""
import base64
import io
import os
import sys
import zipfile

import docx
import openpyxl
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docs_layout as L  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(BASE, 'api', 'assets')

M506 = 'multas de 1 a 60 UTM según el tamaño de la empresa (Art. 506 CT)'

# ------------------------------------------------------------------ Kit Blindaje Laboral
KB_TODOS = [
    ('GUÍA PRÁCTICA DE APLICACIÓN (DT / SUSESO)', 'GUÍA PRÁCTICA DE APLICACIÓN'),
    ('Kit de Blindaje Laboral Pyme 2026 · Instrumento Legal', 'Kit de Blindaje Laboral Pyme 2026 · Modelo de referencia'),
]

KB = {
    '00_INSTRUCCIONES': [
        ("Registro obligatorio: Los anexos de contrato y protocolos deben cargarse en el portal 'Mi DT' de la Dirección del Trabajo (Registro Electrónico Laboral - REL).",
         'Registro: el contrato debe registrarse en el portal Mi DT dentro de 15 días desde su celebración (Art. 9 bis CT); te recomendamos registrar también sus anexos.'),
        ('Para blindar su empresa frente a inspecciones de la Dirección del Trabajo (DT) y evitar sanciones de hasta 60 UTM, ejecute los siguientes pasos en orden prioritario:',
         'Para ordenar la documentación de su empresa ante una fiscalización de la Dirección del Trabajo (DT) y reducir el riesgo de multas, ejecute los siguientes pasos en orden prioritario:'),
        ('2. TABLA RESUMEN DE MÓDULOS Y MULTAS EVITADAS', '2. TABLA RESUMEN DE MÓDULOS Y MULTAS ASOCIADAS'),
        ('MULTA EVITADA', 'MULTA ASOCIADA (REFERENCIAL)'),
        ('Hasta 60 UTM ($4.000.000+)', '1 a 60 UTM según tamaño (Art. 506)'),
        ('Hasta 60 UTM por trabajador', '1 a 60 UTM según tamaño (Art. 506)'),
        ('Hasta 40 UTM', '1 a 60 UTM según tamaño (Art. 506)'),
        ('Cierre de sumarios sin sanción', 'Rebaja si corrige a tiempo (Art. 511)'),
        ('Firmar anexo reduciendo 2 horas el viernes. Prohibido por ley disminuir sueldos.',
         'Firmar anexo con salida anticipada el viernes (42 horas). Prohibido por ley disminuir sueldos.'),
        ('URGENTE: Si no ejercen facultades directivas, regularizarlos a jornada fiscalizada de 42h.',
         'Si en los hechos cumplen horario o tienen supervisión directa, no corresponde excluirlos de jornada: regularizarlos a 42 horas.'),
        ('Firmar pacto Ley 21.220, garantizar 12h de desconexión y registrar en REL de la DT.',
         'Firmar pacto Ley 21.220 y registrarlo en la DT dentro de 15 días; si el teletrabajador distribuye libremente su horario, respetar 12 horas de desconexión.'),
        ('Blindar viáticos de hospedaje y autorizar marcaje de asistencia por GPS/App autorizada.',
         'Documentar los viáticos y usar un sistema de marcaje móvil que cumpla la normativa de la DT.'),
        ('Auditar carpetas antes de la visita del inspector y hacer valer derechos del empleador.',
         'Revisar las carpetas periódicamente y saber cómo actuar ante una fiscalización.'),
        ('Calendario de control para evitar multas inspectivas por vencimiento de plazos legales.',
         'Calendario de control para no olvidar los plazos legales.'),
        ('Plazo 15 días hábiles', '15 días desde la firma'),
        ('Carga obligatoria de contratos y anexos al Registro Electrónico Laboral (REL).',
         'Registro del contrato en Mi DT (Art. 9 bis CT) y, como buena práctica, de sus anexos.'),
        ('Auditoría de libro de asistencia, liquidaciones y pago previsional en Previred.',
         'Revisión de libro de asistencia, liquidaciones y pago previsional en Previred.'),
        ('[  ] AUDITADO', '[  ] REVISADO'),
        ('Regla de oro inspectiva: En caso de fiscalización presencial, la DT no otorga plazos de gracia para contratos no escriturados o protocolos no difundidos.',
         'En una fiscalización presencial, un contrato no escriturado o un protocolo no difundido se constata en el acto: mantenga la carpeta al día.'),
        ('Reconsideración de multas: Ante cualquier sanción notificada, dispone de 15 días hábiles para solicitar la rebaja o sustitución de multa (Art. 511 CT).',
         'Rebaja de multas: si corrige la infracción dentro de 15 días hábiles desde la notificación, puede pedir una rebaja de al menos 50% (80% en micro y pequeñas empresas, Art. 511 CT). Las micro y pequeñas empresas también pueden pedir, dentro de 30 días, sustituir la multa por capacitación (Art. 506 ter CT).'),
        ("Paso 3 (Carga de Anexos 40 Horas): Seleccione el dependiente por RUT, pulse 'Modificar / Anexo' y registre la nueva jornada semanal (42 horas semanales conforme al hito legal 2026).",
         'Paso 3 (Anexos 40 Horas): busque al trabajador por su RUT y registre el anexo o la modificación del contrato con la nueva jornada de 42 horas semanales.'),
        ("Paso 4 (Reglamento Interno y Ley Karin): En la sección 'Reglamento Interno de Orden, Higiene y Seguridad', cargue el archivo PDF del Protocolo Ley Karin como anexo formal obligatorio.",
         'Paso 4 (Reglamento Interno y Ley Karin): si su empresa tiene 10 o más trabajadores permanentes, incorpore el protocolo a su Reglamento Interno y remita copia a la DT y al Ministerio de Salud dentro de 5 días desde su vigencia (Art. 153 CT). Si tiene menos de 10, entregue el protocolo por escrito a cada trabajador al contratarlo (Art. 154 bis CT).'),
        ('Suite de Cumplimiento Normativo para Pymes de Chile', 'Modelos de referencia para pymes de Chile'),
        ('Modelos referenciales tipo conforme a la legislación vigente al año 2026.', 'Basados en la legislación vigente a octubre de 2026.'),
    ],
    '01_Protocolo': [
        ('Exigible por ley a toda empresa en Chile desde el 1 de agosto de 2024 (Ley 21.643 y DS 44 Mintrab).',
         'Exigible por ley a toda empresa en Chile desde el 1 de agosto de 2024 (Ley 21.643, Art. 211-A del Código del Trabajo).'),
        ('Multas gravísimas de hasta 60 UTM según el tamaño de la empresa (Art. 506 del Código del Trabajo).', M506 + '.'),
        ('Todo el año 2026 y subsiguientes. Debe actualizarse semestralmente o tras cada denuncia investigada.',
         'Recomendamos revisarlo al menos una vez al año y después de cada denuncia investigada.'),
        ('el Decreto Supremo N° 44 del Ministerio del Trabajo y Previsión Social, los artículos 2, 153, 154 y 211-A y siguientes del Código del Trabajo',
         'el Decreto Supremo N° 21 de 2024 (reglamento de las investigaciones) y el Decreto Supremo N° 44 sobre gestión preventiva de los riesgos laborales, ambos del Ministerio del Trabajo y Previsión Social, los artículos 2, 153, 154, 154 bis y 211-A y siguientes del Código del Trabajo'),
    ],
    '02_Matriz': [
        ('Esta matriz simplificada cumple con la exigencia inspectiva para micro y pequeñas empresas.',
         'Esta matriz simplificada sirve como base para identificar los riesgos en micro y pequeñas empresas; complétela con la realidad de cada puesto.'),
    ],
    '04_Acta': [
        ('Esta acta es lo primero que exige el fiscalizador de la DT si la persona acude a denunciar a la Inspección.',
         'Esta acta permite acreditar ante la Inspección del Trabajo que las medidas se adoptaron a tiempo.'),
    ],
    '05_Comprobante': [
        ('La falta de acreditación de entrega del protocolo genera multas de hasta 60 UTM según el tamaño de la empresa (Art. 506 CT).',
         'No poder acreditar la entrega del protocolo puede generar ' + M506 + '.'),
        ('firmado por el todos los trabajadores', 'firmado por todos los trabajadores'),
        ('4. La incorporación del presente protocolo como anexo formal e indisoluble del Reglamento Interno de la Empresa.',
         '4. Que el protocolo forma parte del Reglamento Interno de la Empresa o, si la Empresa no está obligada a tenerlo, me fue entregado por escrito al momento de mi contratación (Art. 154 bis del Código del Trabajo).'),
    ],
    '06_Anexo': [
        ('Modalidad Más Solicitada: Reduce 2 horas el día viernes (salida anticipada a las 16:00 o 17:00 hrs) manteniendo 8.5 hrs de lunes a jueves.',
         'Modalidad: 8,5 horas efectivas de lunes a jueves y 8 horas el viernes, con salida anticipada ese día (42 horas semanales).'),
        ('Prohibición Legal Absoluta: El artículo transitorio de la Ley 21.561 prohíbe disminuir la remuneración del dependiente.',
         'Remuneración: la Ley 21.561 establece que la reducción de jornada no puede implicar una disminución de las remuneraciones.'),
        ('Plazo de Registro: Debe subirse al Registro Electrónico Laboral (REL) de la DT en un plazo de 15 días tras su firma.',
         'Registro: recomendamos registrar el anexo en el portal Mi DT dentro de 15 días desde su firma, igual que el contrato (Art. 9 bis CT).'),
        ('Viernes: De 08:30 horas a 16:30 horas, con 60 minutos destinados a colación (salida anticipada en dos horas).',
         'Viernes: De 08:30 horas a 17:30 horas, con 60 minutos destinados a colación (salida anticipada).'),
        ('En estricto cumplimiento del artículo segundo transitorio de la Ley N° 21.561,', 'En cumplimiento de la Ley N° 21.561,'),
    ],
    '07_Anexo': [
        ('Descansos y Domingos: Respeta el artículo 38 del Código del Trabajo respecto a días de descanso compensatorio y 7 domingos anuales.',
         'Descansos y domingos: considera el Art. 38 (al menos dos domingos de descanso al mes) y el Art. 38 bis (siete domingos adicionales al año) del Código del Trabajo.'),
        ('asegurando el otorgamiento de al menos 7 domingos de descanso en el año calendario adicionales a los descansos regulares.',
         'asegurando que al menos dos de los días de descanso de cada mes calendario se otorguen en día domingo (artículo 38 inciso cuarto) y otorgando, adicionalmente, siete días domingo de descanso en cada año de vigencia del contrato (artículo 38 bis), salvo las excepciones legales.'),
    ],
    '08_Anexo': [
        ('el valor mensual de la remuneración líquida y bruta pactada permanece inalterable.', 'el valor mensual de la remuneración pactada permanece inalterable.'),
    ],
    '09_Pacto': [
        ('REGULARIZACIÓN LEGAL VIGENTE: La Ley N° 21.561 restringió drásticamente el Art. 22 inc. 2° a contar del 26 de abril de 2024. Solo pueden quedar excluidos directivos y cargos sin fiscalización superior por la naturaleza de sus labores.',
         'Cuándo usarlo: el Art. 22 inciso 2° excluye de la limitación de jornada a gerentes, administradores, apoderados con facultades de administración y a quienes trabajan sin fiscalización superior inmediata. Si el trabajador, en los hechos, cumple horario o tiene supervisión directa, la exclusión no corresponde.'),
        ('RIESGO DE MULTA Y HORAS EXTRAS: Mantener dependientes indebidamente bajo Art. 22 genera multas de hasta 60 UTM y reclamos de cobro retroactivo de horas extraordinarias.',
         'Riesgo: mantener indebidamente a un trabajador excluido de jornada puede generar multas (Art. 506 CT) y reclamos de horas extraordinarias.'),
        ('SOLUCIÓN INDISPENSABLE: Suscribir este anexo para regularizar al trabajador y adecuarlo a la jornada ordinaria fiscalizada de 42 horas en el marco del hito 2026.',
         'Qué hace este anexo: deja al trabajador en jornada ordinaria de 42 horas, con registro de asistencia.'),
        ('Con motivo de la reforma introducida por la Ley N° 21.561 —cuya restricción legal rige desde el 26 de abril de 2024, limitando dicha exención exclusivamente a directivos con facultades de administración o dependientes sin fiscalización superior inmediata en razón de la naturaleza de sus servicios—, las partes acuerdan',
         'Atendido que las funciones del Trabajador están sujetas a fiscalización superior inmediata y que, por tanto, no concurren los supuestos de dicha norma, las partes acuerdan'),
        ('(reloj biométrico, tarjeta o aplicación móvil debidamente autorizada por la DT)',
         '(reloj control, tarjeta o aplicación móvil que cumpla la normativa de la DT, con alternativa no biométrica)'),
    ],
    '10_Clausula': [
        ('Obligación Legal Art. 64: El empleador tiene prohibición de retener, descontar comisiones bancarias de Transbank o demorar la entrega de propinas.',
         'Art. 64 CT: el empleador debe entregar íntegramente las propinas, no puede hacer descuentos de ninguna naturaleza sobre ellas ni distribuirlas; su distribución corresponde a los trabajadores que las reciben.'),
        ('Plazo de Entrega: Las propinas pagadas con tarjeta deben transferirse al trabajador o repartirse dentro de los plazos legales fijados (máximo semanal o quincenal).',
         'Plazo: las propinas pagadas con tarjeta se pagan en la fecha acordada, que no puede exceder de 7 días hábiles desde que se recibieron del cliente, con copia del comprobante (Art. 64 CT).'),
        ('Multa DT: La apropiación o retención de propinas es sancionada con multas gravísimas de hasta 60 UTM.',
         'Multa: el incumplimiento se sanciona con ' + M506 + '.'),
        ('Las propinas recaudadas a través de medios electrónicos serán totalizadas y distribuidas de manera transparente entre el personal que conforma el turno [DE FORMA SEMANAL / QUINCENAL CADA DÍA VIERNES], emitiéndose una planilla de liquidación de propinas para respaldo y firma del personal.',
         'Las propinas pagadas con tarjeta u otros medios electrónicos serán liquidadas y pagadas por el Empleador a los trabajadores [FRECUENCIA ACORDADA, EJ.: CADA VIERNES], plazo que no podrá exceder de siete (7) días hábiles desde que se recibieron del cliente, entregando copia del comprobante o voucher respectivo. La distribución de las propinas entre el personal corresponde exclusivamente a los trabajadores que las reciben, conforme al artículo 64 del Código del Trabajo.'),
    ],
    '11_Protocolo': [
        ('al organismo administrador del seguro (Mutual de Seguridad / ACHS)', 'al organismo administrador del seguro de la Ley 16.744 (ACHS, Mutual de Seguridad, IST o ISL, según corresponda)'),
        ('y presentará las acciones legales pertinentes.', 'y evaluará las acciones legales pertinentes.'),
    ],
    '12_Pacto': [
        ('Publicación Anticipada: El rol de turnos debe publicarse con al menos 15 días de anticipación.',
         'Publicación anticipada: recomendamos comunicar el rol de turnos con al menos 15 días de anticipación.'),
        ('Garantía 7 Domingos: Al menos 7 días domingo de descanso al año calendario deben ser otorgados a cada trabajador.',
         'Domingos: al menos dos domingos de descanso al mes (Art. 38 inc. 4°) y siete domingos adicionales al año (Art. 38 bis), salvo excepciones legales (contratos de 30 días o menos o jornadas de hasta 20 horas semanales, entre otras).'),
        ('En todo caso, en el lapso de cada año calendario, al menos siete (7) días domingo serán de descanso forzoso para el trabajador, no pudiendo acumularse más de tres turnos dominicales sucesivos.',
         'En todo caso, al menos dos de los días de descanso de cada mes calendario se otorgarán en día domingo (artículo 38 inciso cuarto) y, adicionalmente, el Trabajador gozará de siete días domingo de descanso en cada año de vigencia del contrato (artículo 38 bis), salvo que concurra alguna de las excepciones legales.'),
    ],
    '13_Pacto': [
        ('Carácter Reversible: Cualquiera de las partes puede revocar unilateralmente el teletrabajo con 30 días de aviso previo.',
         'Carácter reversible: si el teletrabajo se pactó después de iniciada la relación laboral, cualquiera de las partes puede volver al trabajo presencial con 30 días de aviso previo.'),
        ('Seguridad Laboral: Requiere verificación de condiciones ergonómicas en el domicilio del trabajador.',
         'Seguridad laboral: el empleador debe informar los riesgos y evaluar las condiciones del puesto de trabajo en el domicilio, conforme a la normativa de seguridad aplicable al teletrabajo.'),
    ],
    '14_Clausula': [
        ('Mandato Legal Art. 152 quáter J: Los trabajadores a distancia o con jornada ordinaria tienen derecho irrenunciable a desconexión digital.',
         'Art. 152 quáter J: el derecho a desconexión aplica a los teletrabajadores que distribuyen libremente su horario o están excluidos de la limitación de jornada. Para los demás, puede pactarse igual como buena práctica.'),
        ('Prohibición: Queda estrictamente prohibido que jefaturas envíen correos o mensajes de WhatsApp en horarios de descanso.',
         'Descansos: el empleador no puede establecer comunicaciones ni formular órdenes en días de descanso, permisos o feriado anual.'),
    ],
    '15_Asignacion': [
        ('Acreditación ante el SII y DT: Debe estipularse con monto razonable ($20.000 a $45.000 CLP mensuales) para no ser reclasificada como sueldo.',
         'Monto: debe ser razonable y guardar relación con el gasto real que compensa; si lo excede, puede ser considerada remuneración.'),
    ],
    '16_Clausula': [
        ('(DICTAMEN DT ORD. N° 2927/58)', '(RESOLUCIÓN EXENTA N° 38 DE 2024 DE LA DT)'),
        ('Aprobación DT: La DT exige que los sistemas móviles cuenten con certificación de cumplimiento técnico y aviso formal al trabajador.',
         'Requisitos DT: los sistemas electrónicos de asistencia deben cumplir la Resolución Exenta N° 38 de 2024 de la DT (entre otros, certificación del sistema y comprobante de cada marcación).'),
        ('aplicación móvil debidamente certificada ante la Dirección del Trabajo.',
         'aplicación móvil que cumple los requisitos de la Resolución Exenta N° 38 de 2024 de la Dirección del Trabajo.'),
    ],
    '17_Clausula': [
        ('Evita Multas SII y DT: El viático mal redactado es reclasificado por fiscalizadores como sueldo encubierto, cobrando cotizaciones retroactivas.',
         'Riesgo de reclasificación: un viático que no corresponde a gastos reales puede ser considerado remuneración, con cotizaciones e impuestos retroactivos.'),
        ('el viático aquí regulado no constituye remuneración y por consiguiente no está afecto',
         'el viático aquí regulado no constituye remuneración y, en la medida que corresponda a gastos razonables, no está afecto'),
    ],
    '18_Checklist': [
        ('CHECKLIST PREVENTIVO DE LOS 10 DOCUMENTOS EXIGIDOS EN UNA INSPECCIÓN DT', 'CHECKLIST PREVENTIVO: 10 DOCUMENTOS QUE SUELE REVISAR UNA INSPECCIÓN DT'),
        ('Uso Preventivo: Audite mensualmente la carpeta de cada trabajador antes de que se presente un fiscalizador.',
         'Uso preventivo: revise mensualmente la carpeta de cada trabajador antes de que se presente un fiscalizador.'),
        ('Sanción: No exhibir contratos o anexos genera multa de hasta 60 UTM por cada trabajador no regularizado.',
         'Sanción: no escriturar el contrato a tiempo se sanciona con multa de 1 a 5 UTM (Art. 9 CT); otras infracciones, con ' + M506 + '.'),
        ('Ley N° 21.643 y DS 44 Mintrab', 'Ley N° 21.643 y Art. 154 bis CT'),
        ('6. Certificado de pago de cotizaciones (Previred F30/1).', '6. Certificado de pago de cotizaciones (Previred) o certificado F30-1 de la DT.'),
        ('DS N° 40 Mintrab Art. 21', 'DS N° 44 Mintrab (reemplazó al DS 40)'),
        ('CERTIFICACIÓN INTERNA: Constancia de Auditoría.', 'CONSTANCIA INTERNA DE REVISIÓN.'),
        ('TRABAJADOR(A) AUDITADO(A)', 'TRABAJADOR(A)'),
    ],
    '19_Guia': [
        ('Clave Legal: Solicitar siempre plazo de 5 a 10 días para exhibir antecedentes complementarios conforme al Art. 506 del Código del Trabajo.',
         'Clave: si falta un documento, pida que quede constancia y que se fije una fecha para presentarlo en la Inspección.'),
        ('a) Solicitar Credencial Oficial: Pida amablemente la credencial de fiscalizador de la DT y verifique su código QR en el sitio oficial de la Dirección del Trabajo.',
         'a) Solicitar la credencial: pida amablemente la credencial del fiscalizador de la DT y verifique su identidad.'),
        ('Derecho a Exhibir Documentos en Soporte Digital: La DT está legalmente obligada a aceptar libros de asistencia, contratos y liquidaciones firmadas digitalmente.',
         'Documentos en soporte digital: puede exhibir contratos, liquidaciones y registros de asistencia electrónicos si su sistema cumple la normativa de la DT.'),
        ('Derecho a Solicitar Plazo para Acompañar Antecedentes: Si falta un documento contable o liquidación antigua, tiene derecho a pedir hasta 5 a 10 días hábiles para presentarlo en la Inspección Comunal antes de que se curse la multa.',
         'Plazo para acompañar antecedentes: si falta un documento, solicite al fiscalizador que lo cite a presentarlo en la Inspección en una fecha determinada y deje constancia de la solicitud.'),
        ('EXIJA escribir al final del acta', 'pida dejar constancia al final del acta'),
        ('Plazo Fatal de 15 Días Hábiles: Desde que le notifican la resolución de multa, tiene exactamente 15 días hábiles administrativos para solicitar la reconsideración administrativa.',
         'Corrija dentro de 15 días hábiles: si corrige la infracción dentro de los 15 días hábiles siguientes a la notificación de la multa, puede pedir su rebaja (Art. 511 CT).'),
        ('Rebaja Sustantiva (de hasta un 80%) o Sustitución por Capacitación: De conformidad al Art. 511 y Art. 506 ter del Código del Trabajo, las micro y pequeñas empresas (hasta 49 trabajadores) tienen el derecho legal de solicitar la reconsideración administrativa de la sanción, accediendo a una rebaja sustantiva de la multa (de hasta un 80% según el tramo de la empresa y gravedad) o su sustitución íntegra por la asistencia obligatoria a un programa de capacitación laboral asistida.',
         'Rebaja o sustitución por capacitación: la rebaja es de al menos 50% de la multa, y de al menos 80% en micro y pequeñas empresas (Art. 511 CT). Además, las micro y pequeñas empresas que no hayan reclamado la multa pueden pedir, dentro de 30 días desde su notificación y por única vez al año por la misma infracción, sustituirla por un programa de capacitación de la DT o, en materias de higiene y seguridad, por un programa de asistencia al cumplimiento (Art. 506 ter CT).'),
    ],
}

# ------------------------------------------------------------------ Kit Ley 21.719
KD_TODOS = [
    ('privacidad@empresa.cl', '[CORREO DE PRIVACIDAD DE LA EMPRESA]'),
    ('https://www.empresa.cl', '[SITIO WEB DE LA EMPRESA]'),
    ('plazo legal fatal de treinta (30) días corridos', 'plazo legal de treinta (30) días corridos'),
]

KD = {
    '0_MANUAL': [
        ('La Ley N° 21.719 y la Dirección del Trabajo (DT) no exigen que una pyme instale costosos servidores ni gaste $2.000.000 en estudios jurídicos. Lo que la ley exige formalmente es el Principio de Responsabilidad y Diligencia Debida:',
         'La Ley N° 21.719 no exige que una pyme instale servidores costosos ni contrate grandes asesorías. Lo que pide es aplicar el principio de responsabilidad:'),
        ('Con esto cumples la Resolución Exenta N° 38 de 2024 de la Dirección del Trabajo sobre relojes control y blindas la confidencialidad de licencias médicas.',
         'Con esto dejas por escrito la alternativa no biométrica que pide la Resolución Exenta N° 38 de 2024 de la Dirección del Trabajo para los relojes control y la reserva de las licencias médicas.'),
        ('Servirán como prueba de cumplimiento si la Dirección del Trabajo o la Agencia de Protección de Datos Personales (APDP) te fiscalizan.',
         'Te sirven para demostrar diligencia si la Agencia de Protección de Datos Personales (APDP) o la Dirección del Trabajo te los piden.'),
        ('La ley es IDÉNTICA para todas las empresas de Chile. El 85% del tratamiento de datos (sueldos, boletas, facturas, contratos y licencias médicas) es exactamente el mismo en cualquier rubro.',
         'La ley es la misma para todas las empresas de Chile, y la mayor parte del tratamiento de datos (sueldos, boletas, facturas, contratos y licencias médicas) es igual en cualquier rubro.'),
        ('firma obligatoriamente el Documento 3 (DPA)', 'firma el Documento 3 (DPA)'),
        ("'Zona videovigilada conforme a la Ley N° 21.719'.", "'Zona videovigilada', indicando dónde consultar la política de privacidad."),
        ('La Dirección del Trabajo (Res. Ex. N° 38/2024) prohíbe despedirlo o sancionarlo. Debes ofrecerle una alternativa no biométrica (tarjeta magnética, clave o libro). Al tener firmado nuestro Documento 1, tu empresa adopta el estándar preventivo exigido por la DT para prevenir sanciones.',
         'No puedes sancionarlo por negarse: la Resolución Exenta N° 38 de 2024 de la DT exige ofrecerle una alternativa no biométrica (tarjeta, clave o libro). El Documento 1 deja esa alternativa por escrito.'),
        ('remitir el reporte a la Agencia sin dilaciones indebidas (meta 72 horas).',
         'remitir el reporte a la Agencia sin dilaciones indebidas (la ley no fija horas; el kit usa 72 horas como meta interna).'),
        ('No. Las micro, pequeñas y medianas empresas están legalmente exentas. Para acreditarlo formalmente ante cualquier inspector, firma el Documento 7 (Test DPO) incluido en este kit.',
         'No es obligatorio: para las empresas privadas, el delegado de protección de datos es voluntario. Solo debes designarlo si adoptas un modelo de prevención de infracciones (también voluntario). El Documento 7 te ayuda a dejar por escrito tu decisión.'),
        ('tienen firmado el anexo de datos personales y huella archivado', 'tienen firmado el anexo de datos personales archivado'),
        ('en el plazo legal fatal de 30 días corridos.', 'en el plazo legal de 30 días corridos (prorrogable una vez).'),
        ('Documento 7 (Test DPO): Test de 4 preguntas firmado por el representante legal que acredita la exención formal de contratar un DPO.',
         'Documento 7 (Guía DPO): decisión sobre el delegado de protección de datos, firmada por el representante legal.'),
    ],
    '1_Anexo': [
        ('(Conforme a la Ley N° 21.719, Arts. 154 bis y ter del Código del Trabajo y Res. Exenta N° 38/2024 de la Dirección del Trabajo)',
         '(Ley N° 19.628 modificada por la Ley N° 21.719, Art. 154 ter del Código del Trabajo y Res. Exenta N° 38/2024 de la Dirección del Trabajo)'),
        ('con fecha 13 de diciembre de 2024 fue promulgada la Ley N° 21.719', 'con fecha 13 de diciembre de 2024 se publicó la Ley N° 21.719'),
        ('en conformidad con los artículos 154 bis y 154 ter del Código del Trabajo', 'en conformidad con el artículo 154 ter del Código del Trabajo'),
        ('En virtud del artículo 12 de la Ley N° 21.719 y el artículo 154 bis del Código del Trabajo,',
         'En virtud de la Ley N° 19.628, modificada por la Ley N° 21.719, y del artículo 154 ter del Código del Trabajo,'),
        ('En cumplimiento estricto del dictamen y Resolución Exenta N° 38 de 2024', 'En cumplimiento de la Resolución Exenta N° 38 de 2024'),
        ('proveedores tecnológicos externos debidamente homologados', 'proveedores externos'),
        ('contado desde la recepción de la solicitud, de conformidad con lo establecido en el artículo 11',
         'contado desde la recepción de la solicitud, prorrogable por una sola vez hasta por otros treinta (30) días corridos, de conformidad con lo establecido en el artículo 11'),
        ('(plazo de prescripción de acciones laborales según los Arts. 480 y 510 del Código del Trabajo)',
         '(plazo que cubre la prescripción de las acciones laborales y previsionales)'),
    ],
    '2_Politica': [
        ('(ARTÍCULO 4 LEY 21.719)', '(LEY N° 19.628, MODIFICADA POR LA LEY N° 21.719)'),
        ('(Art. 13 letra a Ley 21.719)', '(Ley N° 19.628, modificada por la Ley N° 21.719)'),
        ('Conforme a los artículos 5° al 11 de la Ley N° 19.628 (modificada por la Ley N° 21.719)', 'Conforme a la Ley N° 19.628 (modificada por la Ley N° 21.719)'),
        ('contado desde su recepción (Artículo 11 Ley N° 19.628).',
         'contado desde su recepción, prorrogable por una sola vez hasta por otros treinta (Artículo 11 Ley N° 19.628).'),
    ],
    '3_Clausula': [
        ('(Conforme al Artículo 14 bis y 14 ter de la Ley N° 21.719 de Chile)', '(Ley N° 19.628, modificada por la Ley N° 21.719: tratamiento de datos por cuenta del responsable)'),
        ('2. OBLIGACIONES LEGALES DEL ENCARGADO (ARTÍCULO 14 BIS).', '2. OBLIGACIONES DEL ENCARGADO.'),
        ('estará legalmente obligado a notificar al Responsable', 'se obliga a notificar al Responsable'),
    ],
    '6_Formulario': [
        ('reconocidos en los artículos 5° al 11 de la Ley N° 19.628 (modificada por la Ley N° 21.719).', 'reconocidos en la Ley N° 19.628 (modificada por la Ley N° 21.719).'),
        ('PLAZO LEGAL FATAL DE RESPUESTA (ARTÍCULO 11 LEY N° 19.628):', 'PLAZO LEGAL DE RESPUESTA (ARTÍCULO 11 LEY N° 19.628):'),
        ('dentro del plazo fatal de treinta (30) días corridos', 'dentro del plazo de treinta (30) días corridos, prorrogable por una sola vez hasta por otros treinta,'),
    ],
    '7_Guia': [
        ('Criterios de Obligatoriedad y Régimen de Exención para Micro, Pequeñas y Medianas Empresas (MIPYMES)', 'Guía de decisión para micro, pequeñas y medianas empresas (MiPymes)'),
        ('NO, COMO REGLA GENERAL. La Ley N° 21.719 introduce la figura del Delegado de Protección de Datos (DPO u Oficial de Privacidad) inspirada en el estándar internacional (RGPD Art. 37). Sin embargo, el legislador chileno consagró que para las empresas privadas la designación es VOLUNTARIA, salvo que concurra alguna de las causales taxativas de excepción legal que se evalúan a continuación.',
         'NO. La Ley N° 21.719 incorpora la figura del Delegado de Protección de Datos, pero para las empresas privadas su designación es VOLUNTARIA. El delegado forma parte del modelo de prevención de infracciones, que también es voluntario: si tu empresa decide adoptar ese modelo, debe designar un delegado. Las preguntas siguientes te ayudan a dejar por escrito tu decisión.'),
        ('2. TEST DE EVALUACIÓN RÁPIDA DE OBLIGATORIEDAD (4 PREGUNTAS)', '2. PREGUNTAS PARA DECIDIR (4 PREGUNTAS)'),
        ('1. ¿Su entidad es un órgano de la Administración del Estado, Municipalidad o empresa pública?',
         '1. ¿Su empresa adoptará un modelo de prevención de infracciones (voluntario)? Si responde SÍ, debe designar un delegado.'),
        ('2. ¿La actividad principal de su empresa consiste en operaciones que requieren una observación habitual y sistemática masiva de personas (ej: empresas de telecomunicaciones, bancos, burós de crédito)?',
         '2. ¿Su empresa trata datos de muchas personas de forma habitual y sistemática (ej.: seguimiento de clientes, geolocalización)? Si responde SÍ, recomendamos designar un responsable interno.'),
        ('3. ¿La actividad principal de su empresa consiste en el tratamiento masivo y a gran escala de datos sensibles de salud o condenas penales (ej: grandes hospitales, laboratorios farmacéuticos)?',
         '3. ¿Su empresa trata datos sensibles (salud, biometría o datos de menores de edad) como parte habitual de su actividad? Si responde SÍ, recomendamos designar un responsable interno.'),
        ('4. ¿El volumen de titulares tratados en su negocio supera habitualmente el 10% de la población del país?',
         '4. ¿Ha recibido solicitudes de acceso, rectificación o supresión de datos, o reclamos de titulares, en el último año? Si responde SÍ, conviene que una persona coordine las respuestas.'),
        ('[ X ]', '[   ]'),
        ('3. CONCLUSIÓN Y EVALUACIÓN PREVENTIVA', '3. DECISIÓN DE LA EMPRESA'),
        ("Al haber respondido 'NO' a las cuatro preguntas del test, se concluye documentalmente que la empresa NO se encuentra en las causales taxativas de obligatoriedad y no requiere contratar ni designar a un Delegado de Protección de Datos (DPO) externo. Las funciones de coordinación y respuesta ante la Agencia pueden ser ejercidas internamente por la administración o gerencia general sin costo adicional.",
         'Si respondió NO a la pregunta 1, la empresa no está obligada a designar un Delegado de Protección de Datos. Aun así, recomendamos que una persona de la administración o gerencia coordine las respuestas a los titulares y a la Agencia. Responsable interno designado: [NOMBRE Y CARGO].'),
    ],
}

# Planilla RAT (celdas): reemplazos de texto
KD_RAT = [
    ('Inventario Legal Obligatorio exigido por el Artículo 14 ter de la Ley N° 21.719 de Protección de Datos Personales en Chile',
     'Inventario de tratamientos de datos personales (Ley N° 19.628, modificada por la Ley N° 21.719)'),
    ("la nueva Ley N° 21.719 exige llevar un 'Libro de Datos Personales' (el RAT).",
     'la Ley N° 21.719 refuerza tus deberes de información y responsabilidad, y la forma práctica de cumplirlos es llevar un inventario de tus datos personales (el RAT).'),
    ('Es simplemente un inventario oficial donde', 'Es simplemente un inventario donde'),
    ('2. ¿POR QUÉ LA AUTORIDAD (O UN ABOGADO) TE LO VA A PEDIR?', '2. ¿PARA QUÉ SIRVE?'),
    ('Ante cualquier denuncia, inspección de la Dirección del Trabajo (DT) o auditoría de la nueva Agencia de Protección de Datos Personales (APDP), este Excel es el primer documento que te pedirán para comprobar si tu empresa cumple la ley.',
     'Ante un reclamo o una fiscalización de la Agencia de Protección de Datos Personales (APDP), este inventario te permite mostrar rápidamente qué datos tratas, para qué y por cuánto tiempo.'),
    ("No tener el RAT constituye una 'Infracción Grave' que arriesga multas de hasta 10.000 UTM.",
     'Las infracciones a la ley se sancionan con multas de hasta 5.000, 10.000 o 20.000 UTM según sean leves, graves o gravísimas.'),
    ("'📋 Registro RAT Oficial'", "'📋 Registro RAT'"),
    ('Paso 3: ¡Listo! Las 6 filas ya vienen completamente rellenadas con los procesos habituales de una pyme chilena. No necesitas inventar nada. Guárdalo en tu computador o imprímelo para tu carpeta legal.',
     'Paso 3: Revisa las 6 filas: vienen completadas con los procesos habituales de una pyme chilena; ajústalas a tu realidad (agrega o elimina filas). Guárdalo en tu computador o imprímelo para tu carpeta.'),
    ('Conforme al Artículo 14 ter de la Ley N° 21.719 - Agencia de Protección de Datos Personales (APDP) Chile', 'Ley N° 19.628, modificada por la Ley N° 21.719'),
    ('Base Legal de Licitud (Art. 13)', 'Base Legal de Licitud'),
    ('¿Datos Sensibles? (Art. 12)', '¿Datos Sensibles?'),
    ('Medidas de Seguridad Aplicadas (Art. 14 bis)', 'Medidas de Seguridad Aplicadas'),
    ('(Arts. 154 bis y ter)', '(Art. 154 ter)'),
    (' (Arts. 480 y 510 CT)', ''),
    ('Proveedor del reloj control de asistencia homologado y certificado ante la DT', 'Proveedor del sistema de control de asistencia (Res. Exenta N° 38/2024 DT)'),
    (' (Art. 13 letra c)', ''),
    (' (Art. 13 letra e)', ''),
    ('Septiembre de 2026', '[MES Y AÑO]'),
]


# ------------------------------------------------------------------ utilidades
def _paragraphs(d):
    """Todos los párrafos del cuerpo (incluidas tablas), encabezados y pies."""
    yield from (L.Paragraph(p, d) for p in d.element.body.iter(L.W_P))
    # Encabezados y pies que existen en el archivo (sin crear definiciones nuevas)
    for rel in d.part.rels.values():
        if rel.reltype.endswith(('/header', '/footer')):
            yield from (L.Paragraph(p, d) for p in rel.target_part.element.iter(L.W_P))


def _replace_in_paragraph(par, old, new):
    """Reemplaza `old` por `new` aunque esté repartido en varios runs. Conserva el formato del primer run."""
    runs = [r for r in par._p.findall(L.W_R)]
    texts = [''.join(t.text or '' for t in r.findall(L.W_T)) for r in runs]
    full = ''.join(texts)
    i = full.find(old)
    if i < 0:
        return 0
    n = 0
    while i >= 0:
        j = i + len(old)
        pos = 0
        first = None
        for r, t in zip(runs, texts):
            a, b = pos, pos + len(t)
            pos = b
            if b <= i or a >= j or not t:
                continue
            lo, hi = max(i, a) - a, min(j, b) - a
            nodes = r.findall(L.W_T)
            if first is None:
                first = r
                keep = t[:lo] + new + t[hi:]
                rpr = r.find(L.W_RPR)
                if rpr is not None and rpr.find(qn('w:highlight')) is not None:
                    rpr.remove(rpr.find(qn('w:highlight')))
            else:
                keep = t[:lo] + t[hi:]
            nodes[0].text = keep
            nodes[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            for extra in nodes[1:]:
                extra.text = ''
        n += 1
        runs = [r for r in par._p.findall(L.W_R)]
        texts = [''.join(t.text or '' for t in r.findall(L.W_T)) for r in runs]
        full = ''.join(texts)
        i = full.find(old, i + len(new))
    return n


def fix_docx(data, reglas, highlight):
    d = docx.Document(io.BytesIO(data))
    hechos, faltan = 0, []
    for old, new in reglas:
        c = sum(_replace_in_paragraph(p, old, new) for p in _paragraphs(d))
        if c:
            hechos += c
        elif not any(new in p.text for p in _paragraphs(d)):
            faltan.append(old)
    if hechos:
        if highlight == 'kit':
            sect = next((k for k in d.element.body.iterchildren() if k.tag == L.W_P and k.find('.//' + qn('w:sectPr')) is not None), None)
            L.highlight_fields(d, after=sect)
        elif highlight == 'todo':
            L.highlight_fields(d)
        out = io.BytesIO()
        d.save(out)
        data = out.getvalue()
    return data, hechos, faltan


def fix_xlsx(data, reglas):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    hechos, usados = 0, set()
    for ws in wb.worksheets:
        if ws.title == '📋 Registro RAT Oficial':
            ws.title = '📋 Registro RAT'
            hechos += 1
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    v = c.value
                    for old, new in reglas:
                        if old in v:
                            v = v.replace(old, new)
                            usados.add(old)
                    if v != c.value:
                        c.value = v
                        hechos += 1
    textos = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    faltan = [o for o, n in reglas if o not in usados and not (n == '' or any(n in t for t in textos))
              and not (o.startswith("'📋") and '📋 Registro RAT' in [ws.title for ws in wb.worksheets])]
    if hechos:
        out = io.BytesIO()
        wb.save(out)
        data = out.getvalue()
    return data, hechos, faltan


KITS = [
    dict(zip='Kit_Blindaje_Laboral_Pyme_2026.zip', b64='kit-base64.js',
         header='// In-memory base64 package for Vercel Serverless Function (Kit Blindaje Laboral Pyme 2026)\n',
         todos=KB_TODOS, reglas=KB, highlight='kit', sueltos=os.path.join(ASSETS, 'kit_pyme_files')),
    dict(zip='Kit_Ley_21719_Proteccion_Datos_Pyme_2026.zip', b64='kit-datos-base64.js', header='',
         todos=KD_TODOS, reglas=KD, highlight='todo', sueltos=os.path.join(ASSETS, 'kit_datos_files')),
]


def main():
    pendientes = []
    for kit in KITS:
        path = os.path.join(ASSETS, kit['zip'])
        src = zipfile.ZipFile(path)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                data = src.read(info.filename)
                name = info.filename.split('/')[-1]
                if name.endswith('.docx'):
                    clave = next((k for k in kit['reglas'] if name.startswith(k)), None)
                    reglas = kit['todos'] + (kit['reglas'][clave] if clave else [])
                    hl = None if name.startswith(('00_INSTRUCCIONES', '0_MANUAL')) else kit['highlight']
                    data, n, faltan = fix_docx(data, reglas, hl)
                    faltan = [f for f in faltan if (f, ) and f not in [o for o, _ in kit['todos']]]
                    print('  %-70s %3d cambios%s' % (name, n, '  FALTAN: %d' % len(faltan) if faltan else ''))
                    pendientes += [(name, f) for f in faltan]
                elif name.endswith('.pdf') and os.path.exists(os.path.join(kit['sueltos'], name)):
                    # Manual regenerado desde su HTML (scripts/render_kit_manuals.mjs)
                    nuevo = open(os.path.join(kit['sueltos'], name), 'rb').read()
                    print('  %-70s %s' % (name, 'actualizado' if nuevo != data else 'sin cambios'))
                    data = nuevo
                elif name.endswith('.xlsx'):
                    data, n, faltan = fix_xlsx(data, KD_RAT)
                    print('  %-70s %3d cambios%s' % (name, n, '  FALTAN: %d' % len(faltan) if faltan else ''))
                    pendientes += [(name, f) for f in faltan]
                dst.writestr(info, data)
                if kit['sueltos'] and os.path.exists(os.path.join(kit['sueltos'], name)) and not name.endswith('.pdf'):
                    open(os.path.join(kit['sueltos'], name), 'wb').write(data)
        src.close()
        open(path, 'wb').write(buf.getvalue())
        with open(os.path.join(ASSETS, kit['b64']), 'w', encoding='utf-8', newline='\n') as f:
            f.write(kit['header'])
            f.write('module.exports = "%s";\n' % base64.b64encode(buf.getvalue()).decode())
        print(kit['zip'], len(buf.getvalue()), 'bytes')
    if pendientes:
        print('\nTEXTOS NO ENCONTRADOS (ni su versión corregida):')
        for n, f in pendientes:
            print(' -', n, '|', f[:110])


if __name__ == '__main__':
    main()
