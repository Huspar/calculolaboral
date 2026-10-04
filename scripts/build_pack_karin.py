"""Arma el pack gratuito de Ley Karin (lead magnet) a partir de los documentos del Kit Blindaje.

Gratis: estructura del protocolo (para completar), formulario de recepción de denuncia
y comprobante de entrega. El protocolo redactado, la matriz y las actas quedan en el kit.

Uso: python scripts/build_pack_karin.py
"""
import copy
import io
import os
import sys
import zipfile

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docs_layout as L  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KIT = os.path.join(BASE, 'api', 'assets', 'Kit_Blindaje_Laboral_Pyme_2026.zip')
OUT = os.path.join(BASE, 'descargas')
PACK = 'Pack_Ley_Karin_Basico_2026.zip'
FOOTER = 'CálculoLaboral.cl · Pack gratuito Ley Karin 2026 · Modelo de referencia'
HEADER = 'PACK GRATUITO LEY KARIN | CHILE 2026'

FILES = {
    '01_Estructura_Protocolo_Ley_Karin_2026.docx': '01_Protocolo_Integral_Ley_Karin_DS44.docx',
    '02_Formulario_Recepcion_Denuncia_Ley_Karin_2026.docx': '03_Formulario_Recepcion_Denuncia_Karin.docx',
    '03_Comprobante_Entrega_Protocolo_Ley_Karin_2026.docx': '05_Comprobante_Individual_Entrega_Protocolo_Trabajador.docx',
}

GUIA = [
    ('Qué es: ', 'la estructura mínima de un protocolo de prevención según la Ley N° 21.643 y la Circular N° 3.813 de la SUSESO, para completar con los datos de tu empresa.'),
    ('Cómo usarla: ', 'completa cada sección; los campos están resaltados en amarillo. Pide apoyo a tu mutualidad o al ISL para la evaluación de riesgos psicosociales.'),
    ('Versión completa: ', 'el Kit Blindaje Laboral Pyme trae el protocolo ya redactado, la matriz de riesgos y las actas de resguardo y capacitación (calculolaboral.cl/kit-cumplimiento-laboral-pymes).'),
]

SECCIONES = [
    ('PRIMERO: Empresa y ámbito de aplicación. ', 'Identifica a la empresa y a quiénes aplica el protocolo: trabajadores, jefaturas, personal subcontratado y terceros que interactúan con la empresa. Completa: [RAZÓN SOCIAL], RUT [RUT], domicilio [DIRECCIÓN, COMUNA].'),
    ('SEGUNDO: Marco normativo. ', 'Ley N° 21.643 (Ley Karin), artículo 2° y Título IV del Libro II del Código del Trabajo, Decreto Supremo N° 21 de 2024 del Ministerio del Trabajo (procedimiento de investigación), Circular N° 3.813 de la SUSESO y DS N° 44 de 2024 (gestión preventiva de riesgos).'),
    ('TERCERO: Definiciones. ', 'Transcribe las definiciones legales de acoso sexual, acoso laboral y violencia en el trabajo ejercida por terceros (artículo 2° del Código del Trabajo). Recuerda que el acoso laboral puede configurarse con una sola conducta.'),
    ('CUARTO: Identificación de peligros y evaluación de riesgos psicosociales. ', 'Describe cómo evaluarás los riesgos, con perspectiva de género, y quién lo hará. Completa: instrumento usado [CUESTIONARIO CEAL-SM / OTRO], responsable [NOMBRE], fecha de la última evaluación [FECHA].'),
    ('QUINTO: Medidas de prevención y control. ', 'Para cada riesgo identificado, una medida con objetivo medible. Completa: [MEDIDA] · responsable [NOMBRE] · plazo [FECHA] · indicador [CÓMO SE MIDE].'),
    ('SEXTO: Capacitación e información. ', 'Contenidos y frecuencia de las capacitaciones, y cómo informarás cada seis meses los canales de denuncia y las instancias estatales para denunciar. Completa: [FRECUENCIA], [MODALIDAD], [RESPONSABLE].'),
    ('SÉPTIMO: Canal de denuncias. ', 'Quién recibe las denuncias y por qué medios. Las denuncias pueden ser escritas o verbales; si son verbales, se levanta un acta firmada por quien denuncia (usa el formulario de este pack). Completa: receptor [NOMBRE / CARGO], correo [CORREO], otros medios [BUZÓN / FORMULARIO].'),
    ('OCTAVO: Medidas de resguardo. ', 'Al recibir una denuncia se adoptan de inmediato medidas de resguardo (separación de espacios, redistribución de la jornada u otras) y se deriva a la persona denunciante a la atención psicológica temprana del organismo administrador. Completa: organismo administrador [MUTUALIDAD / ISL].'),
    ('NOVENO: Procedimiento de investigación. ', 'Dentro de 3 días hábiles la empresa decide si investiga internamente o remite la denuncia a la Inspección del Trabajo. La investigación interna dura como máximo 30 días hábiles, se realiza por escrito y con confidencialidad, imparcialidad, celeridad y perspectiva de género, y su informe se envía a la Dirección del Trabajo. Completa: investigador designado [NOMBRE / CARGO].'),
    ('DÉCIMO: Medidas y sanciones. ', 'Las medidas o sanciones se aplican según el Reglamento Interno, dentro de 15 días desde que se reciben las observaciones de la Dirección del Trabajo o vence su plazo. Completa: sanciones previstas [AMONESTACIÓN / MULTA / OTRAS].'),
    ('UNDÉCIMO: Privacidad, honra y no represalias. ', 'Resguardo de la privacidad y la honra de todas las personas involucradas, y protección de quien denuncia y de los testigos, cualquiera sea el resultado de la investigación.'),
    ('DUODÉCIMO: Difusión y vigencia. ', 'El protocolo se entrega a cada trabajador (usa el comprobante de este pack) y se incorpora al Reglamento Interno en empresas con 10 o más trabajadores. Completa: fecha de vigencia [FECHA] y periodicidad de revisión [PERIODICIDAD].'),
]


def set_runs(p, *texts):
    """Deja en el párrafo solo los primeros runs con los textos dados, conservando su formato."""
    runs = p.runs
    for i, t in enumerate(texts):
        runs[i].text = t
    for r in runs[len(texts):]:
        r._r.getparent().remove(r._r)


def sect_paragraph(d):
    """Párrafo que cierra la hoja de instrucciones (ver docs_layout.py)."""
    for k in d.element.body.iterchildren():
        if k.tag == L.W_P and k.find('.//' + L.qn('w:sectPr')) is not None:
            return k


def retitle(d, text):
    """Cambia el título que abre el documento limpio (página 2)."""
    t = sect_paragraph(d).getnext()
    node = t.find('.//' + L.W_T)
    node.text = text


def brand(d, title, code):
    t = d.tables[0]
    set_runs(t.cell(0, 0).paragraphs[0], HEADER)
    if title:
        set_runs(t.cell(0, 0).paragraphs[1], title)
    if code:
        set_runs(t.cell(0, 1).paragraphs[0], code)
    for p in d.sections[0].footer.paragraphs:
        if p.runs:
            set_runs(p, FOOTER)


def build_estructura(d):
    TITULO = 'ESTRUCTURA DEL PROTOCOLO DE PREVENCIÓN DEL ACOSO SEXUAL, LABORAL Y LA VIOLENCIA EN EL TRABAJO'
    brand(d, TITULO, 'KARIN-GRATIS-01')
    retitle(d, TITULO)
    cell = d.tables[1].cell(0, 0)
    set_runs(cell.paragraphs[0], 'CÓMO USAR ESTA ESTRUCTURA')
    bullets = cell.paragraphs[1:]
    for p, (pre, txt) in zip(bullets, GUIA):
        set_runs(p, '• ' + pre + txt)
    for p in bullets[len(GUIA):]:
        p._p.getparent().remove(p._p)
    body = [p for p in d.paragraphs if p.runs and p.runs[0].bold and ':' in p.runs[0].text]
    tpl = body[0]._p
    for p in body[1:]:
        p._p.getparent().remove(p._p)
    anchor = tpl
    for i, (pre, txt) in enumerate(SECCIONES):
        el = tpl if i == 0 else copy.deepcopy(tpl)
        if i:
            anchor.addnext(el)
        anchor = el
        set_runs(docx.text.paragraph.Paragraph(el, d), pre, txt)


def fix_texts(d):
    for p in d.paragraphs + [p for t in d.tables for c in t._cells for p in c.paragraphs]:
        for r in p.runs:
            if 'firmado por el todos' in r.text:
                r.text = r.text.replace('firmado por el todos', 'firmado por todos')


def main():
    kit = zipfile.ZipFile(KIT)
    names = {n.split('/')[-1]: n for n in kit.namelist()}
    os.makedirs(OUT, exist_ok=True)
    built = []
    for out_name, src in FILES.items():
        d = docx.Document(io.BytesIO(kit.read(names[src])))
        if out_name.startswith('01_'):
            build_estructura(d)
        else:
            brand(d, None, None)
            set_runs(d.tables[1].cell(0, 0).paragraphs[0], 'GUÍA PRÁCTICA DE APLICACIÓN')
            fix_texts(d)
        path = os.path.join(OUT, out_name)
        L.highlight_fields(d, after=sect_paragraph(d))
        d.save(path)
        built.append(path)
        print('ok', out_name)
    with zipfile.ZipFile(os.path.join(OUT, PACK), 'w', zipfile.ZIP_DEFLATED) as z:
        for path in built:
            z.write(path, arcname=os.path.basename(path))
    print('ok', PACK)


if __name__ == '__main__':
    main()
