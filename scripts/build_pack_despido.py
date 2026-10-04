"""Arma el Pack de Cartas de Despido por Causal (producto de pago) y los modelos gratuitos de la guía.

Formato de cada documento (ver docs_layout.py): la página 1 es la hoja de instrucciones (marca, guía de
uso, lista de revisión) y la carta empieza limpia en la página 2, con los campos [ ... ] resaltados.

Salida:
  api/assets/pack_despido_files/*.docx                       (pack de pago; no se publica)
  api/assets/Pack_Cartas_Despido_por_Causal_Chile_2026.zip
  api/assets/pack-despido-base64.js
  descargas/01_..., 02_..., 03_...docx y descargas/Pack_Modelos_Cartas_Despido_Chile_2026.zip  (gratis)

Uso: python scripts/build_pack_despido.py
"""
import base64
import copy
import os
import shutil
import sys
import zipfile

import docx
from docx.oxml.ns import qn
from docx.table import Table

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docs_layout as L  # noqa: E402
from pack_despido_content import CARTAS, DOCS_EXTRA, CHECKLIST_GRATIS, FINIQ_159, FINIQ_160, COTIZ, GUIA, PLAZO_3  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FREE = os.path.join(BASE, 'descargas')
OUT_DIR = os.path.join(BASE, 'api', 'assets', 'pack_despido_files')
ZIP_PATH = os.path.join(BASE, 'api', 'assets', 'Pack_Cartas_Despido_por_Causal_Chile_2026.zip')
B64_PATH = os.path.join(BASE, 'api', 'assets', 'pack-despido-base64.js')
T_COMPROB = os.path.join(BASE, 'scripts', 'plantillas', 'kit_comprobante.docx')
T_FORM = os.path.join(BASE, 'scripts', 'plantillas', 'kit_formulario.docx')

MARCA_PACK = 'PACK CARTAS DE DESPIDO | CHILE 2026'
PIE_PACK = 'CálculoLaboral.cl · Pack de Cartas de Despido 2026 · Modelo de referencia'
MARCA_FREE = 'MODELO GRATUITO | CHILE 2026'
PIE_FREE = 'CálculoLaboral.cl · Modelo gratuito de carta de despido · Modelo de referencia'
NOTA = ('Versión 2026.10 (octubre de 2026). Modelo de referencia basado en el Código del Trabajo vigente a esa fecha; no constituye '
        'asesoría legal. Si tu caso tiene particularidades (fuero, denuncias de acoso, montos altos), revísalo con un abogado antes de enviar la carta.')
REVISA = [
    'Completé todos los campos resaltados en amarillo y borré las alternativas que no aplican.',
    'Quité el resaltado amarillo del texto final.',
    'Las cotizaciones están pagadas y tengo los comprobantes para adjuntar.',
    'Imprimo desde la página siguiente: esta hoja de instrucciones no se envía.',
    'Guardaré una copia firmada y el comprobante de envío (Correos o recepción en mano).',
]
COMUN = [
    'Completa los campos resaltados en amarillo con datos reales y borra las alternativas que no apliquen.',
    PLAZO_3,
    'Adjunta los comprobantes de pago de cotizaciones al día (AFP, salud y seguro de cesantía). Sin ellos el despido no produce efecto (Art. 162).',
    'En un juicio solo podrás defender los hechos que escribas en la carta (Art. 454 N° 1): descríbelos con fechas, horas, lugares y personas.',
]


def numerar(items):
    return '\n'.join('%d. %s' % (i + 1, t) for i, t in enumerate(items))


class Builder:
    def __init__(self, marca=MARCA_PACK, pie=PIE_PACK):
        self.marca = marca
        self.doc = docx.Document(T_COMPROB)
        body = self.doc.element.body
        kids = [k for k in body.iterchildren() if not k.tag.endswith('sectPr')]
        self.sect = body.find(qn('w:sectPr'))
        self.part = {n: copy.deepcopy(kids[i]) for n, i in dict(
            head=0, rule=1, box=2, spacer=3, body=6, sig=8).items()}
        form = docx.Document(T_FORM)
        fk = [k for k in form.element.body.iterchildren() if not k.tag.endswith('sectPr')]
        self.part['h'] = copy.deepcopy(fk[4])
        # el cuerpo no debe encadenar "mantener con el siguiente"
        for ppr in self.part['body'].iter(qn('w:keepNext')):
            ppr.getparent().remove(ppr)
        for k in kids:
            body.remove(k)
        self.doc.sections[0].footer.paragraphs[0].runs[0].text = pie
        self.sect_p = None

    # --- piezas -------------------------------------------------------------
    def _add(self, el):
        self.sect.addprevious(el)
        return el

    def _para(self, name, *segments):
        el = copy.deepcopy(self.part[name])
        p = docx.text.paragraph.Paragraph(el, self.doc)
        runs = p.runs
        proto = runs[0]._r
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        for i, (text, bold) in enumerate(segments):
            if i == 0:
                r = runs[0]
            else:
                new = copy.deepcopy(proto)
                p._p.append(new)
                r = docx.text.run.Run(new, p)
            r.text = text
            r.bold = bold
        return self._add(el)

    def head(self, titulo, subtitulo, codigo):
        el = copy.deepcopy(self.part['head'])
        t = Table(el, self.doc)
        t.cell(0, 0).paragraphs[0].runs[0].text = self.marca + (' | ' + subtitulo if subtitulo else '')
        t.cell(0, 0).paragraphs[1].runs[0].text = titulo
        t.cell(0, 1).paragraphs[0].runs[0].text = codigo
        self._add(el)
        self._add(copy.deepcopy(self.part['rule']))

    def box(self, titulo, items):
        el = copy.deepcopy(self.part['box'])
        t = Table(el, self.doc)
        cell = t.cell(0, 0)
        paras = cell.paragraphs
        paras[0].runs[0].text = titulo
        proto = paras[1]._p
        for p in paras[1:]:
            p._p.getparent().remove(p._p)
        last = paras[0]._p
        for it in items:
            np_ = copy.deepcopy(proto)
            np = docx.text.paragraph.Paragraph(np_, self.doc)
            for r in np.runs[1:]:
                r._r.getparent().remove(r._r)
            np.runs[0].text = '• ' + it
            last.addnext(np_)
            last = np_
        self._add(el)
        self._add(copy.deepcopy(self.part['spacer']))

    def revisa(self):
        self.heading('ANTES DE ENVIAR, REVISA')
        for c in REVISA:
            self._para('body', ('☐  ' + c, False))

    def nota(self):
        self._para('body', (NOTA, False))

    def page_break(self):
        """Cierra la hoja de instrucciones: de aquí en adelante, documento limpio."""
        self.sect_p = L.section_break_paragraph(self.doc)
        self._add(self.sect_p)

    def title(self, text):
        self._add(L.title_paragraph(text))

    def date(self):
        self._para('body', ('[CIUDAD], [DÍA] de [MES] de [AÑO].', True))

    def addr(self):
        self._para('body', ('Señor(a):\n[NOMBRE COMPLETO DEL TRABAJADOR]\n', True),
                   ('R.U.T.: [RUT DEL TRABAJADOR]\nCargo: [CARGO O FUNCIÓN]\nDomicilio: [DIRECCIÓN REGISTRADA EN EL CONTRATO]\nComuna / Ciudad: [COMUNA / CIUDAD]', False))

    def salut(self):
        self._para('body', ('De nuestra consideración:', True))

    def heading(self, text):
        self._para('h', (text, True))

    def text(self, t):
        self._para('body', (t, False))

    def lead(self, bold, rest):
        self._para('body', (bold, True), (rest, False))

    def close(self, t='Sin otro particular, le saluda atentamente,'):
        self._para('body', (t, False))
        self._add(copy.deepcopy(self.part['spacer']))

    def spacer(self):
        self._add(copy.deepcopy(self.part['spacer']))

    def sig(self, right=None):
        el = copy.deepcopy(self.part['sig'])
        t = Table(el, self.doc)
        left_label, left_det = 'POR LA EMPRESA (REPRESENTANTE LEGAL)', '[NOMBRE REPRESENTANTE LEGAL]\nR.U.T.: [RUT REPRESENTANTE]\n[RAZÓN SOCIAL EMPRESA]'
        if right:
            r_label, _, r_det = right.partition('\n')
        else:
            r_label, r_det = 'ACUSE DE RECIBO / COPIA TRABAJADOR', ('Nombre: [NOMBRE TRABAJADOR]\nR.U.T.: [RUT TRABAJADOR]\n'
                                                                   'Fecha y hora de notificación:\n____ / ____ / ________   ____:____ hrs.')
        for ci, (lab, det) in enumerate([(left_label, left_det), (r_label, r_det)]):
            runs = t.cell(0, ci).paragraphs[0].runs
            runs[0].text = '_____________________________________\n'
            runs[1].text = lab + '\n'
            runs[2].text = det
        self._add(el)

    def finish(self, path):
        if self.sect_p is not None:
            L.unlink_final_footer(self.doc)
            L.highlight_fields(self.doc, after=self.sect_p)
        else:
            L.highlight_fields(self.doc)
        self.doc.save(path)


def codigo(archivo, prefijo='DESPIDO'):
    return '%s-%s' % (prefijo, archivo.split('_')[0])


def build_carta(c, marca=MARCA_PACK, pie=PIE_PACK, prefijo='DESPIDO'):
    b = Builder(marca, pie)
    plazo = c.get('plazo', 3)
    if c.get('pasos_propios'):
        pasos = list(c['pasos_propios'])
    elif plazo == 6:
        pasos = [COMUN[0], 'PLAZO: para esta causal la carta se entrega o envía por carta certificada dentro de los 6 DÍAS HÁBILES siguientes a la '
                           'separación del trabajador, con copia a la Inspección del Trabajo en el mismo plazo.'] + COMUN[2:]
    else:
        pasos = list(COMUN)
    pasos += c.get('box', [])
    b.head(c['titulo'], c['subtitulo'], codigo(c['archivo'], prefijo))
    b.box('INSTRUCCIONES PARA EL EMPLEADOR', pasos)
    b.revisa()
    b.nota()
    b.page_break()
    b.date()
    b.addr()
    b.salut()
    b.text(c['intro'])
    b.heading('I. CAUSAL LEGAL INVOCADA:')
    b.text(c['causal'])
    b.heading(c.get('h2', 'II. HECHOS EN QUE SE FUNDA:'))
    for h in c['hechos']:
        b.text(h)
    b.heading('III. ESTADO DE PAGO DE COTIZACIONES PREVISIONALES:')
    b.text(COTIZ)
    b.heading('IV. FINIQUITO:')
    b.text(c['finiquito'])
    b.close(c.get('cierre', 'Sin otro particular, le saluda atentamente,'))
    b.sig()
    return b


def build_extra(d, marca=MARCA_PACK, pie=PIE_PACK, prefijo='DESPIDO'):
    b = Builder(marca, pie)
    b.head(d['titulo'], '', codigo(d['archivo'], prefijo))
    split = bool(d.get('box'))
    if split:
        b.box(d['box_titulo'], d['box'])
        b.revisa() if d.get('revisa', True) else None
        b.nota()
        b.page_break()
        if d.get('titulo_en_pagina'):
            b.title(d['titulo'])
    elif d.get('aviso'):
        b.box('AVISO', [d['aviso']])
    for kind, *args in d['body']:
        if kind == 'date':
            b.date()
        elif kind == 'addr':
            b.addr()
        elif kind == 'salut':
            b.salut()
        elif kind == 'h':
            b.heading(args[0])
        elif kind == 't':
            b.text(args[0])
        elif kind == 'l':
            b.lead(args[0], args[1])
        elif kind == 'close':
            b.close(*args)
        elif kind == 'sig':
            b.sig(*args)
        elif kind == 'sp':
            b.spacer()
    return b


def cartas_finiquito(c):
    c = dict(c)
    c['finiquito'] = FINIQ_160 if c['tipo'] == '160' else FINIQ_159 if c['tipo'] == '159' else c['finiquito']
    return c


def main():
    if os.path.isdir(OUT_DIR):
        shutil.rmtree(OUT_DIR)
    os.makedirs(OUT_DIR)
    built = []

    g = dict(GUIA)
    g['aviso'] = g['box'][0]
    g['box'] = None
    build_extra(g).finish(os.path.join(OUT_DIR, g['archivo']))
    built.append(g['archivo'])

    for c in CARTAS:
        c = cartas_finiquito(c)
        build_carta(c).finish(os.path.join(OUT_DIR, c['archivo']))
        built.append(c['archivo'])
    for d in DOCS_EXTRA:
        build_extra(d).finish(os.path.join(OUT_DIR, d['archivo']))
        built.append(d['archivo'])

    built.sort()
    with zipfile.ZipFile(ZIP_PATH, 'w', zipfile.ZIP_DEFLATED) as z:
        for name in built:
            z.write(os.path.join(OUT_DIR, name), arcname=name)
    data = open(ZIP_PATH, 'rb').read()
    with open(B64_PATH, 'w', encoding='utf-8', newline='\n') as f:
        f.write('// In-memory base64 package for Vercel Serverless Function (Pack Cartas de Despido por Causal 2026)\n')
        f.write('module.exports = "%s";\n' % base64.b64encode(data).decode())
    print('pack de pago:', len(built), 'documentos', len(data), 'bytes')

    # ---- modelos gratuitos (mismos nombres de archivo que enlaza el sitio) ----
    free_files = []
    for key, nombre in [('01_', '01_Modelo_Carta_Despido_Art161_Necesidades_Empresa_2026.docx'),
                        ('02_', '02_Modelo_Carta_Despido_Art160_N3_Inasistencia_Injustificada_2026.docx')]:
        c = cartas_finiquito(next(x for x in CARTAS if x['archivo'].startswith(key)))
        c['archivo'] = nombre
        build_carta(c, MARCA_FREE, PIE_FREE, 'DESPIDO-GRATIS').finish(os.path.join(FREE, nombre))
        free_files.append(nombre)
    ck = dict(CHECKLIST_GRATIS)
    ck['aviso'] = 'Es una guía de referencia: cada caso puede tener particularidades.'
    build_extra(ck, MARCA_FREE, PIE_FREE, 'DESPIDO-GRATIS').finish(os.path.join(FREE, ck['archivo']))
    free_files.append(ck['archivo'])
    with zipfile.ZipFile(os.path.join(FREE, 'Pack_Modelos_Cartas_Despido_Chile_2026.zip'), 'w', zipfile.ZIP_DEFLATED) as z:
        for name in free_files:
            z.write(os.path.join(FREE, name), arcname=name)
    print('modelos gratuitos:', len(free_files), 'documentos')


if __name__ == '__main__':
    main()
