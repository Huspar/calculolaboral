"""Genera los formatos gratuitos de liquidación de sueldo (Excel y Word).

Salida:
    descargas/formato-liquidacion-de-sueldo.xlsx  (con fórmulas, ejemplo con el ingreso mínimo)
    descargas/formato-liquidacion-de-sueldo.docx  (modelo en blanco para llenar a mano)

Valores 2026: ingreso mínimo $553.553, UTM $71.721, tramos del impuesto único en UTM
(los mismos de js/constants.js) y tasas AFP con comisión.

Uso: python tools/generar_formato_liquidacion.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT_XLSX = ROOT / "descargas" / "formato-liquidacion-de-sueldo.xlsx"
OUT_DOCX = ROOT / "descargas" / "formato-liquidacion-de-sueldo.docx"

FOREST = "00382E"
AMBER = "FFB703"
SLATE = "475569"

IMM = 553553
UTM = 71721
UF = 40975.41
AFP = [
    ("Modelo", 0.1058), ("Uno", 0.1069), ("PlanVital", 0.1116), ("Habitat", 0.1127),
    ("Capital", 0.1144), ("Cuprum", 0.1144), ("Provida", 0.1145),
]
# (desde UTM, factor, rebaja UTM)
TRAMOS = [
    (0, 0, 0), (13.5, 0.04, 0.54), (30, 0.08, 1.74), (50, 0.135, 4.49),
    (70, 0.23, 11.14), (90, 0.304, 17.8), (120, 0.35, 23.32), (310, 0.40, 38.82),
]

MONEY = '"$"#,##0;-"$"#,##0'


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Liquidación"
    par = wb.create_sheet("Parámetros")

    # ---------- Hoja de parámetros ----------
    bold = Font(bold=True)
    par["A1"] = "Parámetros 2026 (actualízalos cuando cambien)"
    par["A1"].font = Font(bold=True, size=12, color=FOREST)
    par["A3"], par["B3"] = "UTM del mes", UTM
    par["A4"], par["B4"] = "UF del último día del mes", UF
    par["A5"], par["B5"] = "Tope imponible AFP y salud (UF)", 89.9
    par["A6"], par["B6"] = "Tope imponible seguro de cesantía (UF)", 135.1
    par["A7"], par["B7"] = "Ingreso mínimo mensual", IMM
    for r in (3, 4, 7):
        par[f"B{r}"].number_format = MONEY
    par["A9"], par["B9"] = "AFP", "Tasa total (10% + comisión)"
    par["A9"].font = par["B9"].font = bold
    for i, (name, rate) in enumerate(AFP, start=10):
        par[f"A{i}"], par[f"B{i}"] = name, rate
        par[f"B{i}"].number_format = "0.00%"
    afp_last = 9 + len(AFP)

    t0 = afp_last + 2
    par[f"A{t0}"], par[f"B{t0}"], par[f"C{t0}"] = "Impuesto único: desde (UTM)", "Factor", "Rebaja (UTM)"
    for c in "ABC":
        par[f"{c}{t0}"].font = bold
    for i, (desde, factor, rebaja) in enumerate(TRAMOS, start=t0 + 1):
        par[f"A{i}"], par[f"B{i}"], par[f"C{i}"] = desde, factor, rebaja
        par[f"B{i}"].number_format = "0.0%"
    tr_first, tr_last = t0 + 1, t0 + len(TRAMOS)
    par.column_dimensions["A"].width = 40
    par.column_dimensions["B"].width = 26
    par.column_dimensions["C"].width = 14
    par.page_setup.fitToWidth = 1
    par.page_setup.fitToHeight = 1
    par.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------- Hoja de liquidación: haberes y descuentos lado a lado ----------
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (31, 15, 2.5, 31, 15)):
        ws.column_dimensions[col].width = w

    ink, muted = "1E293B", "64748B"
    hair = Side(style="thin", color="D7DEE4")
    strong = Side(style="medium", color=FOREST)
    tint = PatternFill("solid", fgColor="E8F0EE")
    soft = PatternFill("solid", fgColor="F1F5F4")
    entry_fill = PatternFill("solid", fgColor="FFFBEA")
    forest_fill = PatternFill("solid", fgColor=FOREST)
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    def put(ref, value, size=9, bold=False, color=ink, fill=None, align=left, fmt=None, border=None):
        c = ws[ref]
        c.value = value
        c.font = Font(name="Arial", size=size, bold=bold, color=color)
        c.alignment = align
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt
        if border:
            c.border = border
        return ref

    # Encabezado
    put("A1", "[Razón social de la empresa]", size=12, bold=True, fill=entry_fill)
    ws.merge_cells("D1:E1")
    put("D1", "LIQUIDACIÓN DE SUELDO", size=15, bold=True, color=FOREST, align=right)
    put("A2", "RUT: ", size=9, color=muted, fill=entry_fill)
    put("D2", "Período", size=9, color=muted, align=right)
    put("E2", "Septiembre 2026", size=9, fill=entry_fill, align=right)
    put("A3", "Dirección: ", size=9, color=muted, fill=entry_fill)
    for col in "ABCDE":
        ws[f"{col}4"].border = Border(bottom=Side(style="thick", color=FOREST))
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[4].height = 6

    # Ficha del trabajador (rótulo | valor, en dos columnas)
    ficha = [
        (("Nombre del trabajador", ""), ("RUT trabajador", "")),
        (("Cargo", ""), ("Fecha de ingreso", "")),
        (("Tipo de contrato", "Indefinido"), ("AFP", "Modelo")),
        (("Sistema de salud", "Fonasa"), ("Plan Isapre en pesos (si aplica)", 0)),
        (("Días trabajados", 30), ("Jornada semanal", "42 horas")),
    ]
    refs = {}
    for i, pair in enumerate(ficha, start=6):
        for (label, val), (lc, vc) in zip(pair, (("A", "B"), ("D", "E"))):
            put(f"{lc}{i}", label, size=8, bold=True, color=muted, fill=soft, border=Border(bottom=hair))
            fmt = MONEY if label.startswith("Plan Isapre") else None
            put(f"{vc}{i}", val, size=9, fill=entry_fill, align=right, fmt=fmt, border=Border(bottom=hair))
            refs[label] = f"{vc}{i}"
        ws.row_dimensions[i].height = 17
    contrato, afp, salud = refs["Tipo de contrato"], refs["AFP"], refs["Sistema de salud"]
    isapre = refs["Plan Isapre en pesos (si aplica)"]

    dv_contrato = DataValidation(type="list", formula1='"Indefinido,Plazo fijo"', allow_blank=False)
    dv_afp = DataValidation(type="list", formula1=f"='Parámetros'!$A$10:$A${afp_last}", allow_blank=False)
    dv_salud = DataValidation(type="list", formula1='"Fonasa,Isapre"', allow_blank=False)
    for dv, ref in ((dv_contrato, contrato), (dv_afp, afp), (dv_salud, salud)):
        ws.add_data_validation(dv)
        dv.add(ref)

    # Tablas de haberes (A:B) y descuentos (D:E)
    def header(row, lc, vc, title):
        b = Border(bottom=strong)
        put(f"{lc}{row}", title.upper(), size=8, bold=True, color=FOREST, fill=tint, border=b)
        put(f"{vc}{row}", "MONTO", size=8, bold=True, color=FOREST, fill=tint, border=b, align=right)

    def item(row, lc, vc, label, value=None, formula=None, entry=False, total=False):
        b = Border(bottom=hair)
        fill = soft if total else None
        put(f"{lc}{row}", label, size=9, bold=total, fill=fill, border=b)
        put(f"{vc}{row}", formula if formula is not None else value, size=9, bold=total,
            fill=entry_fill if entry else fill, border=b, align=right, fmt=MONEY)
        return f"{vc}{row}"

    r0 = 12
    for r in range(r0, r0 + 15):
        ws.row_dimensions[r].height = 17
    # Haberes
    header(r0, "A", "B", "Haberes imponibles")
    base = item(r0 + 1, "A", "B", "Sueldo base", IMM, entry=True)
    item(r0 + 2, "A", "B", "Gratificación", 0, entry=True)
    item(r0 + 3, "A", "B", "Horas extra", 0, entry=True)
    item(r0 + 4, "A", "B", "Comisiones", 0, entry=True)
    item(r0 + 5, "A", "B", "Bonos", 0, entry=True)
    last_imp = item(r0 + 6, "A", "B", "Semana corrida", 0, entry=True)
    imponible = item(r0 + 7, "A", "B", "Total imponible", formula=f"=SUM({base}:{last_imp})", total=True)
    header(r0 + 8, "A", "B", "Haberes no imponibles")
    col = item(r0 + 9, "A", "B", "Colación", 0, entry=True)
    item(r0 + 10, "A", "B", "Movilización", 0, entry=True)
    item(r0 + 11, "A", "B", "Viáticos", 0, entry=True)
    via = item(r0 + 12, "A", "B", "Asignación familiar", 0, entry=True)
    no_imp = item(r0 + 13, "A", "B", "Total no imponible", formula=f"=SUM({col}:{via})", total=True)

    # Descuentos
    tope = "'Parámetros'!$B$5*'Parámetros'!$B$4"
    tope_ces = "'Parámetros'!$B$6*'Parámetros'!$B$4"
    salud7 = f"ROUND(MIN({imponible},{tope})*0.07,0)"
    utm = "'Parámetros'!$B$3"
    tramos_a = f"'Parámetros'!$A${tr_first}:$A${tr_last}"
    tramos_b = f"'Parámetros'!$B${tr_first}:$B${tr_last}"
    tramos_c = f"'Parámetros'!$C${tr_first}:$C${tr_last}"
    header(r0, "D", "E", "Descuentos legales")
    d_afp = item(r0 + 1, "D", "E", "AFP (10% + comisión)",
                 formula=f"=ROUND(MIN({imponible},{tope})*VLOOKUP({afp},'Parámetros'!$A$10:$B${afp_last},2,FALSE),0)")
    d_salud = item(r0 + 2, "D", "E", "Salud (7% o plan Isapre)",
                   formula=f'=IF({salud}="Isapre",MAX({salud7},{isapre}),{salud7})')
    d_ces = item(r0 + 3, "D", "E", "Seguro de cesantía (0,6%)",
                 formula=f'=IF({contrato}="Indefinido",ROUND(MIN({imponible},{tope_ces})*0.006,0),0)')
    rli = f"({imponible}-{d_afp}-MIN({d_salud},{salud7})-{d_ces})"
    d_imp = item(r0 + 4, "D", "E", "Impuesto único",
                 formula=(f"=MAX(0,ROUND({rli}*LOOKUP({rli}/{utm},{tramos_a},{tramos_b})"
                          f"-LOOKUP({rli}/{utm},{tramos_a},{tramos_c})*{utm},0))"))
    legales = item(r0 + 5, "D", "E", "Total descuentos legales", formula=f"=SUM({d_afp}:{d_imp})", total=True)
    header(r0 + 6, "D", "E", "Otros descuentos")
    ant = item(r0 + 7, "D", "E", "Anticipo de sueldo", 0, entry=True)
    item(r0 + 8, "D", "E", "Cuota sindical", 0, entry=True)
    item(r0 + 9, "D", "E", "Préstamo caja de compensación", 0, entry=True)
    otros_last = item(r0 + 10, "D", "E", "Otros descuentos autorizados", 0, entry=True)
    otros = item(r0 + 11, "D", "E", "Total otros descuentos", formula=f"=SUM({ant}:{otros_last})", total=True)

    # Totales generales
    rt = r0 + 15
    box = Border(top=strong, bottom=strong)
    put(f"A{rt}", "TOTAL HABERES", size=10, bold=True, border=box)
    haberes = put(f"B{rt}", f"={imponible}+{no_imp}", size=10, bold=True, border=box, align=right, fmt=MONEY)
    put(f"D{rt}", "TOTAL DESCUENTOS", size=10, bold=True, border=box)
    descuentos = put(f"E{rt}", f"={legales}+{otros}", size=10, bold=True, border=box, align=right, fmt=MONEY)
    ws.row_dimensions[rt].height = 20

    # Bases de cálculo (izquierda) y líquido a pagar (derecha)
    rb = rt + 2
    put(f"A{rb}", "BASES DE CÁLCULO", size=8, bold=True, color=muted)
    bases = [
        ("Base imponible AFP y salud", f"=MIN({imponible},{tope})"),
        ("Base seguro de cesantía", f'=IF({contrato}="Indefinido",MIN({imponible},{tope_ces}),0)'),
        ("Base tributable (impuesto único)", f"={rli}"),
        ("Valor UTM", f"={utm}"),
    ]
    for i, (label, f) in enumerate(bases, start=rb + 1):
        put(f"A{i}", label, size=9, color=ink, border=Border(bottom=hair))
        put(f"B{i}", f, size=9, align=right, fmt=MONEY, border=Border(bottom=hair))
    ws.merge_cells(f"D{rb}:E{rb}")
    put(f"D{rb}", "LÍQUIDO A PAGAR", size=9, bold=True, color=AMBER, fill=forest_fill, align=Alignment(horizontal="left", vertical="center", indent=1))
    ws[f"E{rb}"].fill = forest_fill
    ws.merge_cells(f"D{rb + 1}:E{rb + 3}")
    liquido = put(f"D{rb + 1}", f"={haberes}-{descuentos}", size=22, bold=True, color="FFFFFF",
                  fill=forest_fill, align=Alignment(horizontal="center", vertical="center"), fmt=MONEY)
    for r in range(rb + 1, rb + 5):
        for c in "DE":
            ws[f"{c}{r}"].fill = forest_fill
    ws.merge_cells(f"D{rb + 4}:E{rb + 4}")
    put(f"D{rb + 4}", "Son: ____________________________ pesos", size=8, color="D9E6E3", fill=forest_fill,
        align=Alignment(horizontal="center", vertical="center"))

    # Certificación, firmas y nota
    rc = rb + 7
    ws.merge_cells(f"A{rc}:E{rc}")
    put(f"A{rc}", "Certifico que he recibido de mi empleador, a mi entera satisfacción, el líquido a pagar indicado en esta "
        "liquidación y que no tengo cargo ni cobro alguno que hacer por los conceptos que comprende.",
        size=8, color=muted, align=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[rc].height = 26
    rf = rc + 4
    for lc in ("A", "D"):
        ws.merge_cells(f"{lc}{rf}:{chr(ord(lc) + 1)}{rf}")
        put(f"{lc}{rf}", "Firma y timbre del empleador" if lc == "A" else "Firma del trabajador", size=8, color=muted,
            align=Alignment(horizontal="center"))
        for c in (lc, chr(ord(lc) + 1)):
            ws[f"{c}{rf}"].border = Border(top=Side(style="thin", color=ink))
    rn = rf + 2
    ws.merge_cells(f"A{rn}:E{rn}")
    put(f"A{rn}", "Las celdas amarillas se llenan a mano; el resto se calcula solo. Formato gratuito de "
        "calculolaboral.cl/liquidacion-de-sueldo. Es un modelo de referencia; revisa los valores con tu contrato.",
        size=7, color=muted, align=Alignment(wrap_text=True, horizontal="center"))
    ws.row_dimensions[rn].height = 22

    ws.print_area = f"A1:E{rn}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    wb.save(OUT_XLSX)


# ---------- Word: diseño de una página, estilo de los sistemas de remuneraciones ----------
INK = "1E293B"
MUTED = "64748B"
LINE = "D7DEE4"
SOFT = "F1F5F4"
TINT = "E8F0EE"


def _shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _borders(cell, **sides):
    """sides: top/bottom/left/right = (size_eighths, color) o None para quitar."""
    tc_pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        spec = sides.get(side)
        if spec:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1])
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tc_pr.append(b)


def _no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        b.append(el)
    tbl_pr.append(b)


def _cell_margins(table, top=40, bottom=40, left=90, right=90):
    tbl_pr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tbl_pr.append(m)


def _fixed(table, widths_cm):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def _text(cell, text, size=8.5, bold=False, color=INK, align=None, caps=False, first=True):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.all_caps = caps
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def _gap(doc, pts):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(pts)
    r = p.add_run("")
    r.font.size = Pt(1)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # carta
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    sec.left_margin = sec.right_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(8.5)
    normal.paragraph_format.space_after = Pt(0)
    W = 18.39  # ancho útil en cm
    R = WD_ALIGN_PARAGRAPH.RIGHT
    accent = (12, FOREST)
    hair = (4, LINE)

    # 1. Encabezado: empresa a la izquierda, título y período a la derecha
    t = doc.add_table(rows=1, cols=2)
    _no_table_borders(t); _cell_margins(t, 0, 0, 0, 0); _fixed(t, [10.4, W - 10.4])
    left, right = t.rows[0].cells
    _text(left, "[Razón social de la empresa]", size=11, bold=True)
    _text(left, "RUT: ____________________", size=8.5, color=MUTED, first=False)
    _text(left, "Dirección: ______________________________________", size=8.5, color=MUTED, first=False)
    _text(right, "LIQUIDACIÓN DE SUELDO", size=14, bold=True, color=FOREST, align=R)
    _text(right, "Período: ______________ de 20____", size=9, color=MUTED, align=R, first=False)
    # filete de marca bajo el encabezado
    _gap(doc, 6)
    rule = doc.add_table(rows=1, cols=1)
    _no_table_borders(rule); _cell_margins(rule, 0, 0, 0, 0); _fixed(rule, [W])
    _borders(rule.rows[0].cells[0], top=(16, FOREST))
    _text(rule.rows[0].cells[0], "", size=2)
    _gap(doc, 4)

    # 2. Ficha del trabajador: cuadrícula de 4 columnas con rótulo pequeño sobre el valor
    fields = [
        ("Nombre del trabajador", "RUT", "Cargo", "Fecha de ingreso"),
        ("Tipo de contrato", "AFP", "Sistema de salud", "Días trabajados"),
    ]
    g = doc.add_table(rows=2, cols=4)
    _no_table_borders(g); _cell_margins(g, 90, 110, 110, 90); _fixed(g, [W / 4] * 4)
    for i, row in enumerate(fields):
        for j, label in enumerate(row):
            c = g.rows[i].cells[j]
            _shade(c, SOFT)
            _borders(c, bottom=hair if i == 0 else None, right=hair if j < 3 else None)
            _text(c, label, size=6.5, bold=True, color=MUTED, caps=True)
            _text(c, "", size=9, first=False)
    _gap(doc, 8)

    # 3. Haberes | Descuentos, lado a lado
    haberes = [
        ("H", "Haberes imponibles"),
        ("L", "Sueldo base"), ("L", "Gratificación"), ("L", "Horas extra"),
        ("L", "Comisiones"), ("L", "Bonos"), ("L", "Semana corrida"),
        ("S", "Total imponible"),
        ("H", "Haberes no imponibles"),
        ("L", "Colación"), ("L", "Movilización"), ("L", "Viáticos"), ("L", "Asignación familiar"),
        ("S", "Total no imponible"),
    ]
    descuentos = [
        ("H", "Descuentos legales"),
        ("L", "AFP (10% + comisión)"), ("L", "Salud (7% o plan Isapre)"),
        ("L", "Seguro de cesantía (0,6%)"), ("L", "Impuesto único"),
        ("S", "Total descuentos legales"),
        ("H", "Otros descuentos"),
        ("L", "Anticipo de sueldo"), ("L", "Cuota sindical"), ("L", "Préstamo caja de compensación"),
        ("L", "Otros descuentos autorizados"),
        ("S", "Total otros descuentos"),
        ("E", ""), ("E", ""),
    ]
    n = max(len(haberes), len(descuentos))
    col = [6.35, 2.45, 0.8, 6.35, 2.44]
    m = doc.add_table(rows=n + 1, cols=5)
    _no_table_borders(m); _cell_margins(m, 72, 72, 90, 90); _fixed(m, col)

    def fill(side_rows, c0):
        for i in range(n):
            kind, label = side_rows[i] if i < len(side_rows) else ("E", "")
            a, b = m.rows[i].cells[c0], m.rows[i].cells[c0 + 1]
            if kind == "H":
                for c in (a, b):
                    _shade(c, TINT)
                    _borders(c, bottom=(8, FOREST))
                _text(a, label, size=7, bold=True, color=FOREST, caps=True)
                _text(b, "Monto", size=7, bold=True, color=FOREST, caps=True, align=R)
            elif kind == "L":
                for c in (a, b):
                    _borders(c, bottom=hair)
                _text(a, label)
                _text(b, "$", color=MUTED, align=R)
            elif kind == "S":
                for c in (a, b):
                    _shade(c, SOFT)
                    _borders(c, bottom=hair)
                _text(a, label, bold=True)
                _text(b, "$", bold=True, align=R)
            else:
                _text(a, ""); _text(b, "")

    fill(haberes, 0)
    fill(descuentos, 3)
    for i in range(n + 1):
        _text(m.rows[i].cells[2], "")
    # fila de totales generales
    last = m.rows[n].cells
    for c0, label in ((0, "TOTAL HABERES"), (3, "TOTAL DESCUENTOS")):
        a, b = last[c0], last[c0 + 1]
        for c in (a, b):
            _borders(c, top=accent, bottom=accent)
        _text(a, label, size=9, bold=True)
        _text(b, "$", size=9, bold=True, align=R)
    _gap(doc, 10)

    # 4. Bases de cálculo (izquierda) y líquido a pagar destacado (derecha)
    k = doc.add_table(rows=1, cols=3)
    _no_table_borders(k); _cell_margins(k, 80, 80, 140, 140); _fixed(k, [8.8, 0.8, W - 9.6])
    bases, _, liq = k.rows[0].cells
    _borders(bases, top=hair, bottom=hair, left=hair, right=hair)
    _text(bases, "Bases de cálculo", size=6.5, bold=True, color=MUTED, caps=True)
    for line in ("Base imponible AFP y salud:  $", "Base seguro de cesantía:  $", "Base tributable (impuesto único):  $", "Valor UF: $                    Valor UTM: $"):
        _text(bases, line, size=8, color=INK, first=False).paragraph_format.space_before = Pt(3)
    _shade(liq, FOREST)
    _text(liq, "Líquido a pagar", size=7, bold=True, color=AMBER, caps=True)
    _text(liq, "$", size=18, bold=True, color="FFFFFF", first=False).paragraph_format.space_before = Pt(4)
    _text(liq, "Son: _______________________________ pesos", size=7.5, color="D9E6E3", first=False).paragraph_format.space_before = Pt(6)
    _text(k.rows[0].cells[1], "")
    _gap(doc, 8)

    # 5. Forma de pago
    f = doc.add_table(rows=1, cols=3)
    _no_table_borders(f); _cell_margins(f, 50, 50, 0, 90); _fixed(f, [W / 3] * 3)
    for c, (label, val) in zip(f.rows[0].cells, (("Forma de pago", "Transferencia  /  Efectivo  /  Cheque"), ("Banco y cuenta", "______________________"), ("Fecha de pago", "____ / ____ / ________"))):
        _borders(c, bottom=hair)
        _text(c, label, size=6.5, bold=True, color=MUTED, caps=True)
        _text(c, val, size=8.5, first=False)
    _gap(doc, 10)

    # 6. Certificación y firmas
    cert = doc.add_paragraph()
    cert.paragraph_format.space_after = Pt(0)
    r = cert.add_run(
        "Certifico que he recibido de mi empleador, a mi entera satisfacción, el líquido a pagar indicado en esta "
        "liquidación y que no tengo cargo ni cobro alguno que hacer por los conceptos que comprende."
    )
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    _gap(doc, 44)
    s = doc.add_table(rows=1, cols=3)
    _no_table_borders(s); _cell_margins(s, 40, 0, 0, 0); _fixed(s, [7.6, W - 15.2, 7.6])
    for c, label in ((s.rows[0].cells[0], "Firma y timbre del empleador"), (s.rows[0].cells[2], "Firma del trabajador")):
        _borders(c, top=(6, INK))
        _text(c, label, size=7.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
    _text(s.rows[0].cells[1], "")
    _gap(doc, 14)

    n2 = doc.add_paragraph()
    n2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = n2.add_run(
        "Formato gratuito de calculolaboral.cl/liquidacion-de-sueldo. Es un modelo de referencia: los montos dependen "
        "del contrato, la AFP y el sistema de salud de cada trabajador."
    )
    r.font.size = Pt(6.5)
    r.italic = True
    r.font.color.rgb = RGBColor.from_string(MUTED)

    doc.core_properties.title = "Formato de liquidación de sueldo"
    doc.core_properties.author = "Cálculo Laboral"
    doc.save(OUT_DOCX)

if __name__ == "__main__":
    build_xlsx()
    build_docx()
    print(f"OK: {OUT_XLSX.relative_to(ROOT)} y {OUT_DOCX.relative_to(ROOT)}")
